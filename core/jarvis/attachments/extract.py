"""Turn files into text the LLM can read."""
from __future__ import annotations

import io
from pathlib import Path

EXT_KIND = {
    **{e: "image" for e in (".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".heic", ".heif", ".tif", ".tiff")},
    ".pdf": "pdf",
    **{e: "office" for e in (".docx", ".xlsx", ".xlsm", ".pptx")},
    **{e: "audio" for e in (".mp3", ".wav", ".m4a", ".ogg", ".opus", ".flac", ".aac", ".webm", ".amr", ".3gp")},
    **{e: "video" for e in (".mp4", ".mkv", ".mov", ".avi", ".wmv")},
    **{e: "text" for e in (".txt", ".md", ".csv", ".json", ".xml", ".html", ".htm", ".py", ".js", ".ts", ".tsx",
                           ".jsx", ".css", ".java", ".c", ".cpp", ".h", ".cs", ".go", ".rs", ".rb", ".php", ".sh",
                           ".ps1", ".bat", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".log", ".sql", ".kt", ".swift",
                           ".lua", ".srt", ".vtt", ".tex")},
}


def kind_of(name: str, mime: str) -> str:
    ext = Path(name).suffix.lower()
    if ext in EXT_KIND:
        k = EXT_KIND[ext]
        if ext == ".webm" and mime.startswith("video"):
            return "video"
        return k
    for prefix, k in (("image/", "image"), ("audio/", "audio"), ("video/", "video"), ("text/", "text")):
        if mime.startswith(prefix):
            return k
    if mime == "application/pdf":
        return "pdf"
    return "other"


def decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1250", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def looks_text(data: bytes) -> bool:
    sample = data[:4096]
    return bool(sample) and b"\x00" not in sample


def ocr_bytes(data: bytes) -> str | None:
    """Local OCR via rapidocr; None when unavailable."""
    try:
        from rapidocr_onnxruntime import RapidOCR
    except Exception:
        return None
    global _OCR
    try:
        _OCR
    except NameError:
        _OCR = RapidOCR()
    res, _ = _OCR(data)
    return "\n".join(r[1] for r in res or [])


def extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages):
        txt = (page.extract_text() or "").strip()
        if not txt:
            ocr_parts = []
            for img in getattr(page, "images", [])[:4]:
                t = ocr_bytes(img.data)
                if t:
                    ocr_parts.append(t)
            txt = "\n".join(ocr_parts) if ocr_parts else "[stránka bez textu]"
        pages.append(f"--- strana {i + 1} ---\n{txt}")
    return "\n".join(pages)


def extract_docx(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def extract_xlsx(path: Path, max_rows: int = 2000) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"=== list {ws.title} ===")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                out.append("…")
                break
            if any(c is not None for c in row):
                out.append(" | ".join("" if c is None else str(c) for c in row))
    return "\n".join(out)


def extract_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"--- snímek {i} ---")
        for sh in slide.shapes:
            if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
                out.append(sh.text_frame.text)
    return "\n".join(out)


def extract(path: Path, kind: str, transcribe=None) -> str:
    ext = path.suffix.lower()
    if kind == "pdf":
        return extract_pdf(path)
    if kind == "office":
        return {".docx": extract_docx, ".pptx": extract_pptx}.get(ext, extract_xlsx)(path)
    if kind == "text":
        return decode_text(path.read_bytes())
    if kind == "audio":
        if transcribe is None:
            raise RuntimeError("přepis audia není k dispozici (hlasový modul není nainstalovaný)")
        return transcribe(path)
    if kind == "image":
        return ""  # images go to vision models directly; OCR lazily when needed
    if kind == "video":
        return ""
    data = path.read_bytes()
    if looks_text(data):
        return decode_text(data)
    return ""


def make_thumb(path: Path, out: Path, size: int = 320) -> bool:
    try:
        from PIL import Image

        try:
            import pillow_heif  # type: ignore

            pillow_heif.register_heif_opener()
        except Exception:
            pass
        with Image.open(path) as im:
            im = im.convert("RGB")
            im.thumbnail((size, size))
            im.save(out, "JPEG", quality=80)
        return True
    except Exception:
        return False


def to_jpeg_for_llm(path: Path, max_side: int = 1600) -> bytes:
    """Downscale big photos before sending to a vision model."""
    from PIL import Image

    with Image.open(path) as im:
        im = im.convert("RGB")
        im.thumbnail((max_side, max_side))
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=85)
        return buf.getvalue()
