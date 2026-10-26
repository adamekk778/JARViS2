"""OMEGA §16 / §35 running a small business: a register of quotes (numbers, validity, accepted / rejected, win rate),
billing basis from a work log, customer requests with priority and due dates, comparing supplier offers (weighted,
with a check whether the winner holds when the weights move), campaign results (CTR, CPA, ROAS, significance,
where to move the budget), and income / expenses from a bank statement (categories, recurring payments, spikes)."""
from __future__ import annotations

import datetime as dt
import math
import re
import statistics
import time
from collections import Counter, defaultdict
from pathlib import Path

from .business import _col, _kc, _parse_date, quote_totals
from .comms import fold
from .science import column_nums, read_csv

QUOTES_SQL = """CREATE TABLE IF NOT EXISTS quotes (id INTEGER PRIMARY KEY AUTOINCREMENT, number TEXT UNIQUE, client TEXT, created REAL,
  valid_until TEXT, base REAL, total REAL, status TEXT, path TEXT, decided REAL, note TEXT);"""
REQ_SQL = """CREATE TABLE IF NOT EXISTS customer_requests (id INTEGER PRIMARY KEY AUTOINCREMENT, created REAL, client TEXT, text TEXT,
  source TEXT, priority TEXT, status TEXT, due TEXT, closed REAL);"""


def _today(today: str | dt.date | None) -> dt.date:
    if isinstance(today, dt.date):
        return today
    return dt.date.fromisoformat(today) if today else dt.date.today()


# ============================================================================================== 168 quotes register
def next_quote_number(rt, today: dt.date | None = None) -> str:
    rt.db.conn.executescript(QUOTES_SQL)
    y = (today or dt.date.today()).year
    row = rt.db.one("SELECT number FROM quotes WHERE number LIKE ? ORDER BY number DESC LIMIT 1", (f"N-{y}-%",))
    n = int(row["number"].rsplit("-", 1)[1]) + 1 if row else 1
    return f"N-{y}-{n:04d}"


def register_quote(rt, number: str, client: str, base: float, total: float, valid_days: int, path: str, note: str = "",
                   today: dt.date | None = None) -> dict:
    rt.db.conn.executescript(QUOTES_SQL)
    d = today or dt.date.today()
    rt.db.exec("INSERT INTO quotes(number, client, created, valid_until, base, total, status, path, note) VALUES (?,?,?,?,?,?,?,?,?)",
               (number, client, time.mktime(d.timetuple()), (d + dt.timedelta(days=valid_days)).isoformat(), base, total, "otevřená", path, note))
    return {"number": number, "valid_until": (d + dt.timedelta(days=valid_days)).isoformat()}


STATUS = {"prijata": "přijatá", "prijato": "přijatá", "accepted": "přijatá", "ano": "přijatá", "vyhrana": "přijatá",
          "odmitnuta": "odmítnutá", "odmitnuto": "odmítnutá", "rejected": "odmítnutá", "ne": "odmítnutá", "prohrana": "odmítnutá",
          "otevrena": "otevřená", "zrusena": "zrušená"}


def set_quote_status(rt, which: str, status: str) -> dict:
    """`which` = the quote number or the client's name (their newest open quote)."""
    rt.db.conn.executescript(QUOTES_SQL)
    st = STATUS.get(fold(status).replace(" ", ""))
    if not st:
        raise ValueError("Stav: přijatá, odmítnutá, otevřená nebo zrušená.")
    row = rt.db.one("SELECT * FROM quotes WHERE number=?", (which.strip(),))
    if not row:
        f = fold(which)
        rows = [r for r in rt.db.all("SELECT * FROM quotes WHERE status='otevřená' ORDER BY created DESC") if f in fold(r["client"])]
        row = rows[0] if rows else None
    if not row:
        raise ValueError(f"Nabídku „{which}“ neznám (číslo N-RRRR-NNNN nebo jméno klienta s otevřenou nabídkou).")
    rt.db.exec("UPDATE quotes SET status=?, decided=? WHERE id=?", (st, time.time() if st in ("přijatá", "odmítnutá") else None, row["id"]))
    return {"number": row["number"], "client": row["client"], "status": st, "total": row["total"]}


def quotes_overview(rt, today: str | dt.date | None = None, soon_days: int = 7) -> dict:
    rt.db.conn.executescript(QUOTES_SQL)
    d = _today(today)
    rt.db.exec("UPDATE quotes SET status='propadlá' WHERE status='otevřená' AND valid_until < ?", (d.isoformat(),))
    rows = rt.db.all("SELECT * FROM quotes ORDER BY created DESC")
    open_ = [r for r in rows if r["status"] == "otevřená"]
    won = [r for r in rows if r["status"] == "přijatá"]
    lost = [r for r in rows if r["status"] == "odmítnutá"]
    expiring = [{"number": r["number"], "client": r["client"], "total": r["total"], "days": (dt.date.fromisoformat(r["valid_until"]) - d).days}
                for r in open_ if (dt.date.fromisoformat(r["valid_until"]) - d).days <= soon_days]
    days_to_decide = [(r["decided"] - r["created"]) / 86400 for r in won + lost if r["decided"]]
    by_client: dict[str, dict] = defaultdict(lambda: {"quotes": 0, "won": 0, "won_value": 0.0})
    for r in rows:
        c = by_client[r["client"]]
        c["quotes"] += 1
        if r["status"] == "přijatá":
            c["won"] += 1
            c["won_value"] += r["total"]
    decided = len(won) + len(lost)
    return {"open": len(open_), "open_value": round(sum(r["total"] for r in open_), 2), "expiring": sorted(expiring, key=lambda x: x["days"]),
            "expired": sum(1 for r in rows if r["status"] == "propadlá"), "won": len(won), "won_value": round(sum(r["total"] for r in won), 2),
            "lost": len(lost), "win_rate": round(100 * len(won) / decided, 1) if decided else None,
            "avg_days_to_decision": round(statistics.mean(days_to_decide), 1) if days_to_decide else None,
            "avg_won": round(statistics.mean([r["total"] for r in won]), 2) if won else None,
            "clients": dict(sorted(by_client.items(), key=lambda kv: -kv[1]["won_value"])),
            "recent": [{k: r[k] for k in ("number", "client", "total", "status", "valid_until")} for r in rows[:10]]}


# ============================================================================================== 169 billing basis
def billing(path: Path, client: str = "", period: str = "", default_rate: float = 0, vat_pct: float = 21, vat_payer: bool = True) -> dict:
    """Work log CSV (date, client, project, hours, rate, description, expenses, billed?) → per client: hours × rate
    by project, re-billed expenses, totals with VAT, what was done; only unbilled rows when a "billed" column exists."""
    header, rows, _, _ = read_csv(path)
    di, ci = _col(header, "datum", "date"), _col(header, "klient", "client", "zákazník", "zakaznik", "odběratel", "customer")
    pi, hi = _col(header, "projekt", "project", "zakázka", "zakazka"), _col(header, "hodiny", "hours", "odpracováno")
    ri, ti = _col(header, "sazba", "rate", "kč/h"), _col(header, "popis", "description", "činnost", "cinnost", "práce", "prace")
    ei, bi = _col(header, "náklady", "naklady", "výdaje", "vydaje", "expenses", "materiál", "material"), _col(header, "fakturováno", "fakturovano", "billed", "vyfakturováno")
    if ci is None or hi is None:
        raise ValueError("Potřebuju sloupce klient a hodiny (a ideálně datum, projekt, sazba, popis).")

    def col(i):
        return column_nums([r[i] if i is not None and i < len(r) else "" for r in rows]) if i is not None else [None] * len(rows)
    hours, rates, exps = col(hi), col(ri), col(ei)
    groups: dict[str, dict] = {}
    skipped_billed = 0
    for k, r in enumerate(rows):
        get = (lambda i: r[i].strip() if i is not None and i < len(r) else "")  # noqa: E731
        cl = get(ci)
        if not cl or (client and fold(client) not in fold(cl)):
            continue
        d = _parse_date(get(di)) if di is not None else None
        if period and (not d or not d.isoformat().startswith(period)):
            continue
        if bi is not None and fold(get(bi)) in ("ano", "yes", "1", "x", "true", "ok", "fakturovano"):
            skipped_billed += 1
            continue
        g = groups.setdefault(cl, {"projects": defaultdict(lambda: {"hours": 0.0, "amount": 0.0, "rates": set(), "what": Counter()}),
                                   "expenses": 0.0, "days": set(), "missing_rate": 0})
        h = hours[k] or 0.0
        rate = rates[k] if rates[k] else default_rate
        if not rate:
            g["missing_rate"] += 1
        p = g["projects"][get(pi) or "práce"]
        p["hours"] += h
        p["amount"] += h * (rate or 0)
        p["rates"].add(rate or 0)
        if get(ti):
            p["what"][get(ti)] += 1
        g["expenses"] += exps[k] or 0.0
        if d:
            g["days"].add(d.isoformat())
    out = {}
    for cl, g in groups.items():
        items = []
        for name, p in g["projects"].items():
            rate = round(p["amount"] / p["hours"], 2) if p["hours"] else 0
            items.append({"name": f"{name} ({', '.join(w for w, _ in p['what'].most_common(3))})" if p["what"] else name, "qty": round(p["hours"], 2),
                          "unit": "h", "price": rate})
        if g["expenses"]:
            items.append({"name": "Přefakturované náklady", "qty": 1, "unit": "ks", "price": round(g["expenses"], 2)})
        t = quote_totals(items, vat_pct, 0, vat_payer)
        days = sorted(g["days"])
        out[cl] = {"items": t["lines"], "base": t["base"], "vat": t["vat"], "total": t["total"], "hours": round(sum(p["hours"] for p in g["projects"].values()), 2),
                   "period": f"{days[0]} – {days[-1]}" if days else period, "missing_rate": g["missing_rate"]}
    return {"clients": out, "skipped_billed": skipped_billed, "total": round(sum(c["total"] for c in out.values()), 2)}


def billing_docx(dest: Path, client: str, b: dict, supplier: str) -> Path:
    import docx

    d = docx.Document()
    d.add_heading(f"Podklad k fakturaci – {client}", 0)
    d.add_paragraph(f"Dodavatel: {supplier}\nObdobí: {b['period']}\nOdpracováno: {b['hours']:g} h")
    tab = d.add_table(rows=1, cols=5)
    tab.style = "Light Grid Accent 1"
    for c, h in zip(tab.rows[0].cells, ["Položka", "Množství", "Sazba", "Základ", "Celkem s DPH"]):
        c.text = h
    for ln in b["items"]:
        for c, v in zip(tab.add_row().cells, [ln["name"], f"{ln['qty']:g} {ln['unit']}", _kc(ln["price"]), _kc(ln["base"]), _kc(ln["total"])]):
            c.text = v
    p = d.add_paragraph()
    p.add_run(f"Celkem bez DPH: {_kc(b['base'])}\nDPH: {_kc(b['vat'])}\nCelkem: {_kc(b['total'])}").bold = True
    d.add_paragraph("Podklad, ne daňový doklad – fakturu vystav ve svém fakturačním programu.")
    d.save(str(dest))
    return dest


# ============================================================================================== 170 customer requests
URGENT = ("nefunguje", "nejde", "vypadek", "havarie", "spadl", "urgent", "hned", "okamzite", "kriticke", "blokuje", "chyba v produkci", "asap")
LOW = ("az budete mit cas", "nespecha", "nekdy", "do budoucna", "napad", "low priority")
DUE_DAYS = {"vysoká": 1, "normální": 5, "nízká": 14}


def priority_of(text: str) -> str:
    f = fold(text)
    if any(k in f for k in URGENT):
        return "vysoká"
    if any(k in f for k in LOW):
        return "nízká"
    return "normální"


def add_request(rt, client: str, text: str, source: str = "", priority: str = "", due: str = "", today: dt.date | None = None) -> dict:
    rt.db.conn.executescript(REQ_SQL)
    pr = priority or priority_of(text)
    if pr not in DUE_DAYS:
        raise ValueError("Priorita: vysoká, normální nebo nízká.")
    d = today or dt.date.today()
    due = due or (d + dt.timedelta(days=DUE_DAYS[pr])).isoformat()
    rt.db.exec("INSERT INTO customer_requests(created, client, text, source, priority, status, due) VALUES (?,?,?,?,?,?,?)",
               (time.mktime(d.timetuple()), client, text, source, pr, "otevřený", due))
    rid = rt.db.one("SELECT max(id) AS id FROM customer_requests")["id"]
    return {"id": rid, "client": client, "priority": pr, "due": due}


def requests(rt, client: str = "", include_closed: bool = False, today: str | dt.date | None = None) -> dict:
    rt.db.conn.executescript(REQ_SQL)
    d = _today(today)
    rows = rt.db.all("SELECT * FROM customer_requests" + ("" if include_closed else " WHERE status='otevřený'") + " ORDER BY id")
    out = []
    for r in rows:
        if client and fold(client) not in fold(r["client"]):
            continue
        late = (d - dt.date.fromisoformat(r["due"])).days if r["status"] == "otevřený" else None
        out.append({"id": r["id"], "client": r["client"], "text": r["text"], "priority": r["priority"], "status": r["status"], "due": r["due"],
                    "overdue_days": late if late and late > 0 else 0, "age_days": (d - dt.date.fromtimestamp(r["created"])).days})
    order = {"vysoká": 0, "normální": 1, "nízká": 2}
    out.sort(key=lambda x: (x["status"] != "otevřený", -x["overdue_days"], order.get(x["priority"], 1), x["due"]))
    per = Counter(x["client"] for x in out if x["status"] == "otevřený")
    return {"requests": out, "open": sum(1 for x in out if x["status"] == "otevřený"), "overdue": sum(1 for x in out if x["overdue_days"]),
            "by_client": dict(per.most_common())}


def close_request(rt, which: str) -> dict:
    rt.db.conn.executescript(REQ_SQL)
    row = rt.db.one("SELECT * FROM customer_requests WHERE id=? AND status='otevřený'", (int(which),)) if which.strip().isdigit() else None
    if not row:
        f = fold(which)
        cands = [r for r in rt.db.all("SELECT * FROM customer_requests WHERE status='otevřený' ORDER BY id DESC")
                 if f in fold(r["text"]) or f in fold(r["client"])]
        row = cands[0] if cands else None
    if not row:
        raise ValueError(f"Otevřený požadavek „{which}“ nevidím.")
    rt.db.exec("UPDATE customer_requests SET status='vyřízený', closed=? WHERE id=?", (time.time(), row["id"]))
    return {"id": row["id"], "client": row["client"], "text": row["text"]}


# ============================================================================================== 172 supplier offers
CRITERIA = {"price": ("cena", False), "delivery_days": ("dodání", False), "warranty_months": ("záruka", True),
            "payment_days": ("splatnost", True), "rating": ("hodnocení", True), "tco": ("náklady za dobu užívání", False)}
DEFAULT_W = {"price": 0.45, "delivery_days": 0.15, "warranty_months": 0.15, "payment_days": 0.1, "rating": 0.15}


def _scores(offers: list[dict], weights: dict[str, float]) -> list[dict]:
    res = []
    crit = [c for c in weights if weights[c] > 0 and any(o.get(c) is not None for o in offers)]
    rng = {}
    for c in crit:
        vals = [float(o[c]) for o in offers if o.get(c) is not None]
        rng[c] = (min(vals), max(vals))
    for o in offers:
        num = den = 0.0
        parts, missing = {}, []
        for c in crit:
            if o.get(c) is None:
                missing.append(CRITERIA[c][0])
                continue
            lo, hi = rng[c]
            n = 1.0 if hi == lo else (float(o[c]) - lo) / (hi - lo)
            if not CRITERIA[c][1]:
                n = 1 - n
            parts[c] = round(n, 3)
            num += weights[c] * n
            den += weights[c]
        res.append({"supplier": o["supplier"], "score": round(100 * num / den, 1) if den else 0.0, "parts": parts, "missing": missing})
    return sorted(res, key=lambda x: -x["score"])


def compare_offers(offers: list[dict], weights: dict[str, float] | None = None, years: float = 0) -> dict:
    """Offers [{supplier, price, delivery_days, warranty_months, payment_days, rating, yearly_cost?, shipping?}] →
    weighted score (each criterion scaled best = 1, worst = 0), ranking, and whether the winner survives each weight
    being halved or raised by half."""
    if len(offers) < 2:
        raise ValueError("Na srovnání potřebuju aspoň dvě nabídky.")
    offers = [dict(o) for o in offers]
    w = dict(weights or DEFAULT_W)
    unknown = [k for k in w if k not in CRITERIA]
    if unknown:
        raise ValueError(f"Neznámá kritéria: {', '.join(unknown)}. Umím: {', '.join(CRITERIA)}.")
    if years:
        for o in offers:
            o["tco"] = float(o.get("price") or 0) + float(o.get("shipping") or 0) + float(o.get("yearly_cost") or 0) * years
        w["tco"] = w.pop("price", 0.45)
    ranking = _scores(offers, w)
    winner = ranking[0]["supplier"]
    flips = []
    for c in w:
        for f in (0.5, 1.5):
            alt = _scores(offers, {**w, c: w[c] * f})[0]["supplier"]
            if alt != winner:
                flips.append(f"{CRITERIA[c][0]} ×{f:g} → vyhraje {alt}")
    return {"ranking": ranking, "winner": winner, "margin": round(ranking[0]["score"] - ranking[1]["score"], 1), "robust": not flips, "flips": flips,
            "weights": {CRITERIA[k][0]: round(v / sum(w.values()), 3) for k, v in w.items()},
            "tco": {o["supplier"]: round(o["tco"], 2) for o in offers} if years else None}


# ============================================================================================== 173 campaigns
def _z_p(c1: float, n1: float, c2: float, n2: float) -> float:
    if not n1 or not n2:
        return 1.0
    p = (c1 + c2) / (n1 + n2)
    se = math.sqrt(p * (1 - p) * (1 / n1 + 1 / n2))
    if se == 0:
        return 1.0
    z = (c1 / n1 - c2 / n2) / se
    return 2 * (1 - statistics.NormalDist().cdf(abs(z)))


def campaigns(header: list[str], rows: list[list[str]]) -> dict:
    """Campaign table (campaign, spend, impressions, clicks, conversions, revenue[, date]) → per campaign CTR, CPC, CR,
    CPA, ROAS, ROI; whether its conversion rate really differs from the rest (z-test); where to move the budget."""
    ki = _col(header, "kampaň", "kampan", "campaign", "název", "nazev", "name")
    si = _col(header, "náklady", "naklady", "spend", "cost", "útrata", "utrata", "výdaje")
    ii, cli = _col(header, "zobrazení", "zobrazeni", "impressions", "impr"), _col(header, "prokliky", "kliky", "clicks", "kliknutí")
    coi, ri = _col(header, "konverze", "conversions", "nákupy", "objednávky", "objednavky"), _col(header, "tržby", "trzby", "revenue", "hodnota", "value", "příjem")
    if ki is None or si is None:
        raise ValueError("Potřebuju sloupce kampaň a náklady (a ideálně zobrazení, prokliky, konverze, tržby).")

    def nums(i):
        return column_nums([r[i] if i is not None and i < len(r) else "" for r in rows]) if i is not None else [None] * len(rows)
    sp, im, cl, co, rv = nums(si), nums(ii), nums(cli), nums(coi), nums(ri)
    agg: dict[str, dict] = defaultdict(lambda: {"spend": 0.0, "impressions": 0.0, "clicks": 0.0, "conversions": 0.0, "revenue": 0.0})
    for k, r in enumerate(rows):
        name = r[ki].strip() if ki < len(r) else ""
        if not name:
            continue
        a = agg[name]
        for key, arr in (("spend", sp), ("impressions", im), ("clicks", cl), ("conversions", co), ("revenue", rv)):
            a[key] += arr[k] or 0.0
    tot = {k: sum(a[k] for a in agg.values()) for k in ("spend", "impressions", "clicks", "conversions", "revenue")}

    def metrics(a):
        return {"ctr_pct": round(100 * a["clicks"] / a["impressions"], 2) if a["impressions"] else None,
                "cpc": round(a["spend"] / a["clicks"], 2) if a["clicks"] else None,
                "cr_pct": round(100 * a["conversions"] / a["clicks"], 2) if a["clicks"] else None,
                "cpa": round(a["spend"] / a["conversions"], 2) if a["conversions"] else None,
                "roas": round(a["revenue"] / a["spend"], 2) if a["spend"] else None,
                "roi_pct": round(100 * (a["revenue"] - a["spend"]) / a["spend"], 1) if a["spend"] else None}
    out = []
    for name, a in agg.items():
        m = {"campaign": name, **{k: round(v, 2) for k, v in a.items()}, **metrics(a)}
        if a["clicks"] and tot["clicks"] - a["clicks"] > 0:
            p = _z_p(a["conversions"], a["clicks"], tot["conversions"] - a["conversions"], tot["clicks"] - a["clicks"])
            m["cr_differs"] = p < 0.05
            m["p_value"] = round(p, 4)
        out.append(m)
    out.sort(key=lambda m: -(m["roas"] or 0))
    losers = [m for m in out if m["roas"] is not None and m["roas"] < 1]
    best = out[0] if out and out[0]["roas"] else None
    move = None
    if losers and best and best["roas"] >= 1:
        freed = sum(m["spend"] for m in losers)
        move = {"from": [m["campaign"] for m in losers], "to": best["campaign"], "amount": round(freed, 2),
                "extra_revenue_if_roas_holds": round(freed * best["roas"] - sum(m["revenue"] for m in losers), 2)}
    return {"campaigns": out, "total": {**{k: round(v, 2) for k, v in tot.items()}, **metrics(tot)}, "move_budget": move}


# ============================================================================================== 174 income and expenses
DEFAULT_RULES = {
    "bydlení": ["najem", "najemne", "svj", "hypoteka", "energie", "cez", "pre ", "e.on", "eon", "innogy", "plyn", "voda", "teplo", "elektrina"],
    "jídlo": ["albert", "lidl", "tesco", "kaufland", "billa", "penny", "globus", "rohlik", "kosik", "potraviny", "makro"],
    "restaurace": ["restaurace", "bistro", "kavarna", "wolt", "bolt food", "dame jidlo", "damejidlo", "foodora", "mcdonald", "kfc", "starbucks"],
    "doprava": ["benzina", "shell", "omv", "mol ", "orlen", "cerpaci", "dpp", "jizdenka", "cd ", "regiojet", "leo express", "uber", "bolt", "parkovani"],
    "předplatné": ["netflix", "spotify", "youtube", "apple.com", "google", "microsoft", "disney", "hbo", "voyo", "o2 tv", "adobe", "chatgpt", "openai", "anthropic"],
    "telefon a internet": ["vodafone", "t-mobile", "o2 czech", "o2 ", "upc", "internet", "nordic telecom"],
    "zdraví": ["lekarna", "dr.max", "benu", "lekar", "nemocnice", "zubar", "pojistovna"],
    "nákupy": ["alza", "mall", "datart", "ikea", "notino", "zalando", "aboutyou", "temu", "aliexpress", "amazon", "hornbach", "obi"],
    "pojištění": ["pojisteni", "kooperativa", "allianz", "generali", "csob pojistovna", "uniqa"],
    "hotovost": ["vyber", "bankomat", "atm"],
    "poplatky": ["poplatek", "vedeni uctu", "urok z uveru", "sankce"],
}
INCOME_RULES = {"mzda": ["mzda", "plat ", "vyplata", "salary"], "faktury": ["faktura", "vs ", "uhrada fa"], "vratky": ["vratka", "refund", "dobropis"]}


def categorize(text: str, amount: float, rules: dict[str, list[str]] | None = None) -> str:
    f = " " + fold(text) + " "
    for cat, keys in (rules or {}).items():
        if any(fold(k) in f for k in keys):
            return cat
    for cat, keys in (DEFAULT_RULES if amount < 0 else INCOME_RULES).items():
        if any(k in f for k in keys):
            return cat
    return "ostatní" if amount < 0 else "ostatní příjmy"


def _who(text: str) -> str:
    t = re.sub(r"\d{4,}|\d+[.,]\d+|\b\d{1,2}[./]\d{1,2}[./]?\d{0,4}\b", " ", fold(text))
    return " ".join(t.split()[:3])


def cashflow(header: list[str], rows: list[list[str]], rules: dict[str, list[str]] | None = None) -> dict:
    """Bank statement (Fio, ČSOB, KB, Air, Revolut… exports) → months (income, expenses, net, savings rate), expenses by
    category, counterparties, recurring payments (subscriptions), unusual spikes and the largest single payments."""
    di = _col(header, "datum zaúčtování", "datum provedení", "datum", "date", "booking date", "completed date", "started date")
    ai = _col(header, "částka", "castka", "objem", "amount", "suma")
    if di is None or ai is None:
        raise ValueError("Ve výpisu nevidím sloupce datum a částka.")
    texts = [i for i, h in enumerate(header) if i not in (di, ai) and any(k in fold(h) for k in (
        "protiucet", "nazev", "zprava", "poznamka", "popis", "description", "merchant", "obchodnik", "typ", "identifikace", "counterparty", "ucel"))]
    amounts = column_nums([r[ai] if ai < len(r) else "" for r in rows])
    tx = []
    for r, a in zip(rows, amounts):
        d = _parse_date(r[di]) if di < len(r) else None
        if d is None or a is None:
            continue
        text = " ".join(r[i] for i in texts if i < len(r) and r[i].strip())
        tx.append({"date": d, "amount": a, "text": text, "category": categorize(text, a, rules), "who": _who(text)})
    if not tx:
        raise ValueError("Žádné řádky s datem a částkou.")
    months: dict[str, dict] = defaultdict(lambda: {"income": 0.0, "expenses": 0.0})
    cats: dict[str, float] = defaultdict(float)
    cat_month: dict[tuple[str, str], float] = defaultdict(float)
    for t in tx:
        m = t["date"].strftime("%Y-%m")
        if t["amount"] >= 0:
            months[m]["income"] += t["amount"]
        else:
            months[m]["expenses"] += -t["amount"]
            cats[t["category"]] += -t["amount"]
            cat_month[(t["category"], m)] += -t["amount"]
    mon = [{"month": m, "income": round(v["income"], 2), "expenses": round(v["expenses"], 2), "net": round(v["income"] - v["expenses"], 2),
            "savings_pct": round(100 * (v["income"] - v["expenses"]) / v["income"], 1) if v["income"] else None} for m, v in sorted(months.items())]
    recurring = []
    by_who: dict[str, list[dict]] = defaultdict(list)
    for t in tx:
        if t["amount"] < 0 and t["who"]:
            by_who[t["who"]].append(t)
    for who, ts in by_who.items():
        ms = {t["date"].strftime("%Y-%m") for t in ts}
        amts = [-t["amount"] for t in ts]
        if len(ms) >= 3 and len(ts) <= len(ms) + 1 and statistics.pstdev(amts) <= 0.15 * statistics.mean(amts):
            recurring.append({"who": ts[-1]["text"][:60] or who, "monthly": round(statistics.mean(amts), 2), "months": len(ms), "category": ts[-1]["category"]})
    recurring.sort(key=lambda x: -x["monthly"])
    spikes = []
    mlist = [m["month"] for m in mon]
    for cat in cats:
        series = [cat_month.get((cat, m), 0.0) for m in mlist]
        if len(series) >= 3:
            base = statistics.median(series[:-1])
            if (base > 0 and series[-1] > 1.5 * base and series[-1] - base > 500) or (base == 0 and series[-1] > 2000):
                spikes.append({"category": cat, "last": round(series[-1], 2), "usual": round(base, 2)})
    exp_total = sum(cats.values())
    return {"months": mon, "categories": [{"category": c, "amount": round(v, 2), "pct": round(100 * v / exp_total, 1)} for c, v in sorted(cats.items(), key=lambda kv: -kv[1])],
            "recurring": recurring, "recurring_monthly": round(sum(r["monthly"] for r in recurring), 2), "spikes": spikes,
            "largest": [{"date": t["date"].isoformat(), "amount": t["amount"], "text": t["text"][:80]} for t in sorted(tx, key=lambda t: t["amount"])[:5] if t["amount"] < 0],
            "uncategorized_pct": round(100 * cats.get("ostatní", 0) / exp_total, 1) if exp_total else 0, "transactions": len(tx)}
