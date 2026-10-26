"""OMEGA §35 business: customer segments (RFM) from orders, lead scoring by rules you can see, the sales funnel and where
it leaks, a price quote with VAT as a Word document, pricing scenarios with elasticity, a revenue forecast with a
back-test, customer feedback (sentiment and topics) and project profitability. All from the user's own files."""
from __future__ import annotations

import datetime as dt
import math
import re
import statistics
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from .comms import fold
from .science import _date, _num, read_csv


def _col(header: list[str], *names: str) -> int | None:
    h = [fold(x) for x in header]
    for n in names:
        for i, x in enumerate(h):
            if x == fold(n) or fold(n) in x:
                return i
    return None


def _parse_date(s: str) -> dt.date | None:
    d = _date(s)
    if d:
        return d[1]
    try:
        return dt.datetime.fromisoformat(s.strip()[:19]).date()
    except ValueError:
        return None


# ============================================================================================== 437 segmentation
SEGMENTS = [  # (name, test on r/f/m scores 1–5, what to do)
    ("šampioni", lambda r, f, m: r >= 4 and f >= 4 and m >= 4, "odměnit, požádat o recenzi a doporučení"),
    ("věrní", lambda r, f, m: f >= 4 and r >= 3, "nabídnout věrnostní program / vyšší variantu"),
    ("noví", lambda r, f, m: r >= 4 and f <= 2, "dobře přivítat, druhý nákup podpořit"),
    ("slibní", lambda r, f, m: r >= 3 and f <= 3 and m >= 3, "připomenout se, doporučit související zboží"),
    ("ohrožení", lambda r, f, m: r <= 2 and f >= 3, "osobně oslovit dřív, než odejdou"),
    ("usínající", lambda r, f, m: r == 3 and f <= 3, "připomenout se nabídkou"),
    ("ztracení", lambda r, f, m: r <= 2 and f <= 2, "jedna reaktivační nabídka, jinak nechat být"),
]


def _quint(values: dict[str, float], reverse: bool = False) -> dict[str, int]:
    """Scores 1–5 by rank; equal values share one score (the average rank of the tie)."""
    order = sorted(values, key=lambda k: values[k], reverse=reverse)
    n = len(order)
    out, i = {}, 0
    while i < n:
        j = i
        while j + 1 < n and values[order[j + 1]] == values[order[i]]:
            j += 1
        sc = min(5, 1 + int(5 * ((i + j) / 2) / n))
        for k in order[i:j + 1]:
            out[k] = sc
        i = j + 1
    return out


def rfm(path: Path, today: dt.date | None = None) -> dict:
    """Orders CSV (customer, date, amount) → recency / frequency / money scores and segments with an action each."""
    header, rows, _, _ = read_csv(path)
    ci, di, ai = _col(header, "zákazník", "customer", "klient", "email", "jméno"), _col(header, "datum", "date"), \
        _col(header, "částka", "amount", "cena", "celkem", "total", "tržba")
    if None in (ci, di, ai):
        raise ValueError("Potřebuju sloupce zákazník, datum a částka (nebo customer, date, amount).")
    last: dict[str, dt.date] = {}
    freq: Counter = Counter()
    money: dict[str, float] = defaultdict(float)
    bad = 0
    for r in rows:
        try:
            c, d, a = r[ci].strip(), _parse_date(r[di]), _num(r[ai])
        except IndexError:
            bad += 1
            continue
        if not c or d is None or a is None:
            bad += 1
            continue
        last[c] = max(last.get(c, d), d)
        freq[c] += 1
        money[c] += a
    if len(last) < 5:
        raise ValueError("Na segmentaci je potřeba aspoň 5 zákazníků.")
    today = today or max(last.values())
    rec = {c: (today - d).days for c, d in last.items()}
    rs, fs, ms = _quint(rec, reverse=True), _quint(dict(freq)), _quint(dict(money))
    customers = []
    for c in last:
        r_, f_, m_ = rs[c], fs[c], ms[c]
        seg = next((s for s in SEGMENTS if s[1](r_, f_, m_)), ("ostatní", None, "sledovat"))
        customers.append({"customer": c, "days_since": rec[c], "orders": freq[c], "money": round(money[c], 2), "r": r_, "f": f_, "m": m_,
                          "segment": seg[0], "action": seg[2]})
    segs = defaultdict(lambda: {"customers": 0, "money": 0.0})
    for x in customers:
        segs[x["segment"]]["customers"] += 1
        segs[x["segment"]]["money"] += x["money"]
    total = sum(money.values())
    summary = [{"segment": k, "customers": v["customers"], "money": round(v["money"], 2), "share_pct": round(100 * v["money"] / total, 1) if total else 0,
                "action": next((s[2] for s in SEGMENTS if s[0] == k), "sledovat")} for k, v in sorted(segs.items(), key=lambda kv: -kv[1]["money"])]
    return {"customers": sorted(customers, key=lambda x: -x["money"]), "segments": summary, "skipped_rows": bad, "as_of": today.isoformat()}


# ============================================================================================== 438 leads
DEFAULT_RULES = {"rozpočet": 30, "rozhodovatel": 20, "termín do 3 měsíců": 20, "velikost firmy": 15, "zapojení": 15}


def score_leads(path: Path, weights: dict[str, float] | None = None) -> dict:
    """Lead CSV → a 0–100 score from visible rules: budget, decision maker, timeline, company size, engagement."""
    header, rows, _, _ = read_csv(path)
    w = {**DEFAULT_RULES, **(weights or {})}
    ni = _col(header, "jméno", "name", "firma", "company", "lead")
    bi, ri, ti = _col(header, "rozpočet", "budget"), _col(header, "role", "pozice", "title"), _col(header, "termín", "timeline", "kdy")
    si, ei = _col(header, "zaměstnanci", "employees", "velikost", "size"), _col(header, "návštěvy", "visits", "otevření", "opens", "interakce", "engagement")
    if ni is None:
        raise ValueError("Potřebuju aspoň sloupec se jménem / firmou leadu.")
    out = []
    for r in rows:
        get = (lambda i: r[i].strip() if i is not None and i < len(r) else "")  # noqa: E731
        why, score = [], 0.0
        b = _num(get(bi)) if get(bi) else None
        if b:
            part = min(1.0, b / 100000)
            score += w["rozpočet"] * part
            why.append(f"rozpočet {b:,.0f}".replace(",", " "))
        role = fold(get(ri))
        if re.search(r"ceo|cto|cfo|reditel|majitel|jednatel|owner|founder|head|vedouci|manager", role):
            score += w["rozhodovatel"]
            why.append("rozhoduje")
        t = fold(get(ti))
        months = _num(re.sub(r"[^\d.,]", "", t)) if re.search(r"\d", t) else None
        if re.search(r"hned|ihned|asap|tento mesic|now", t) or (months is not None and months <= 3):
            score += w["termín do 3 měsíců"]
            why.append("brzy")
        size = _num(get(si)) if get(si) else None
        if size:
            score += w["velikost firmy"] * min(1.0, math.log10(max(size, 1)) / 3)
        eng = _num(get(ei)) if get(ei) else None
        if eng:
            score += w["zapojení"] * min(1.0, eng / 10)
            if eng >= 5:
                why.append("aktivní")
        total = sum(w.values())
        s = round(100 * score / total) if total else 0
        out.append({"lead": get(ni), "score": s, "grade": "horký" if s >= 70 else "teplý" if s >= 40 else "studený", "why": why})
    out.sort(key=lambda x: -x["score"])
    return {"leads": out, "weights": w, "hot": sum(1 for x in out if x["grade"] == "horký")}


# ============================================================================================== 439 funnel
def funnel(stages: list[tuple[str, float]]) -> dict:
    """Stage counts in order → conversion between steps and overall, the biggest leak, and what +10 % at the leak gives."""
    if len(stages) < 2:
        raise ValueError("Potřebuju aspoň dvě fáze trychtýře.")
    steps = []
    for (a, na), (b, nb) in zip(stages, stages[1:]):
        if nb > na:
            raise ValueError(f"Fáze {b} ({nb:g}) má víc než předchozí {a} ({na:g}) – zkontroluj pořadí.")
        steps.append({"from": a, "to": b, "rate_pct": round(100 * nb / na, 1) if na else 0.0, "lost": na - nb})
    worst = min(steps, key=lambda s: s["rate_pct"])
    overall = stages[-1][1] / stages[0][1] if stages[0][1] else 0
    gain = overall * (min(100, worst["rate_pct"] * 1.1) / worst["rate_pct"]) if worst["rate_pct"] else overall
    return {"steps": steps, "overall_pct": round(100 * overall, 2), "leak": worst,
            "if_leak_plus10": round(stages[0][1] * gain - stages[-1][1], 1)}


def funnel_from_events(path: Path, order: list[str]) -> list[tuple[str, float]]:
    """Events CSV (user, stage) → unique users that reached each stage in the given order (counted only if they
    passed all earlier stages)."""
    header, rows, _, _ = read_csv(path)
    ui, si = _col(header, "uživatel", "user", "zákazník", "id"), _col(header, "fáze", "stage", "událost", "event", "krok")
    if ui is None or si is None:
        raise ValueError("Potřebuju sloupce uživatel a fáze (user, stage).")
    seen: dict[str, set] = defaultdict(set)
    for r in rows:
        if max(ui, si) < len(r):
            seen[r[ui].strip()].add(fold(r[si].strip()))
    out = []
    for i, st in enumerate(order):
        need = [fold(s) for s in order[: i + 1]]
        out.append((st, float(sum(1 for s in seen.values() if all(n in s for n in need)))))
    return out


# ============================================================================================== 441 quote
def quote_totals(items: list[dict], vat_pct: float = 21, discount_pct: float = 0, vat_payer: bool = True) -> dict:
    lines = []
    for it in items:
        qty, price = float(it.get("qty", 1)), float(it["price"])
        vat = float(it.get("vat", vat_pct)) if vat_payer else 0.0
        base = round(qty * price * (1 - discount_pct / 100), 2)
        lines.append({"name": it["name"], "qty": qty, "unit": it.get("unit", "ks"), "price": price, "base": base, "vat_pct": vat,
                      "vat": round(base * vat / 100, 2), "total": round(base * (1 + vat / 100), 2)})
    by_rate: dict[float, dict] = defaultdict(lambda: {"base": 0.0, "vat": 0.0})
    for ln in lines:
        by_rate[ln["vat_pct"]]["base"] += ln["base"]
        by_rate[ln["vat_pct"]]["vat"] += ln["vat"]
    base, vat = round(sum(x["base"] for x in lines), 2), round(sum(x["vat"] for x in lines), 2)
    return {"lines": lines, "base": base, "vat": vat, "total": round(base + vat, 2), "by_rate": {k: {kk: round(vv, 2) for kk, vv in v.items()} for k, v in by_rate.items()},
            "discount_pct": discount_pct, "vat_payer": vat_payer}


def _kc(x: float) -> str:
    return f"{x:,.2f} Kč".replace(",", " ").replace(".", ",")


def quote_docx(dest: Path, t: dict, client: str, supplier: str, title: str = "Cenová nabídka", valid_days: int = 30, note: str = "",
               number: str = "") -> Path:
    import docx
    from docx.shared import Pt

    d = docx.Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    d.add_heading(title + (f" č. {number}" if number else ""), 0)
    today = dt.date.today()
    d.add_paragraph(f"Dodavatel: {supplier}\nOdběratel: {client}\nDatum: {today.day}. {today.month}. {today.year}\n"
                    f"Platnost nabídky: {valid_days} dní (do {(today + dt.timedelta(days=valid_days)).strftime('%d. %m. %Y')})")
    tab = d.add_table(rows=1, cols=6 if t["vat_payer"] else 4)
    tab.style = "Light Grid Accent 1"
    heads = ["Položka", "Množství", "Cena/j.", "Základ", "DPH", "Celkem"] if t["vat_payer"] else ["Položka", "Množství", "Cena/j.", "Celkem"]
    for c, h in zip(tab.rows[0].cells, heads):
        c.text = h
    for ln in t["lines"]:
        cells = tab.add_row().cells
        vals = [ln["name"], f"{ln['qty']:g} {ln['unit']}", _kc(ln["price"]), _kc(ln["base"]), f"{ln['vat_pct']:g} %", _kc(ln["total"])] if t["vat_payer"] \
            else [ln["name"], f"{ln['qty']:g} {ln['unit']}", _kc(ln["price"]), _kc(ln["base"])]
        for c, v in zip(cells, vals):
            c.text = v
    if t["discount_pct"]:
        d.add_paragraph(f"Ceny zahrnují slevu {t['discount_pct']:g} %.")
    if t["vat_payer"]:
        for rate, v in sorted(t["by_rate"].items()):
            d.add_paragraph(f"Základ {rate:g} %: {_kc(v['base'])}, DPH {_kc(v['vat'])}")
        p = d.add_paragraph()
        p.add_run(f"Celkem bez DPH: {_kc(t['base'])}\nDPH: {_kc(t['vat'])}\nCelkem k úhradě: {_kc(t['total'])}").bold = True
    else:
        p = d.add_paragraph()
        p.add_run(f"Celkem: {_kc(t['total'])}").bold = True
        d.add_paragraph("Dodavatel není plátcem DPH.")
    if note:
        d.add_paragraph(note)
    d.save(str(dest))
    return dest


# ============================================================================================== 442 pricing
def price_scenarios(price: float, volume: float, unit_cost: float, fixed_cost: float = 0, elasticity: float = -1.5,
                    changes_pct: list[float] | None = None) -> dict:
    """Constant-elasticity demand: volume' = volume · (price'/price)^elasticity. Revenue, profit and margin for each
    price change, the profit-maximising price and the break-even volume."""
    if price <= 0 or volume < 0:
        raise ValueError("Cena musí být kladná a objem nezáporný.")
    rows = []
    for ch in changes_pct or [-20, -10, -5, 0, 5, 10, 20]:
        p = price * (1 + ch / 100)
        v = volume * (p / price) ** elasticity
        rev = p * v
        prof = (p - unit_cost) * v - fixed_cost
        rows.append({"change_pct": ch, "price": round(p, 2), "volume": round(v, 1), "revenue": round(rev), "profit": round(prof),
                     "margin_pct": round(100 * (p - unit_cost) / p, 1)})
    best = None
    if elasticity < -1:
        p_opt = unit_cost * elasticity / (1 + elasticity)
        v_opt = volume * (p_opt / price) ** elasticity
        best = {"price": round(p_opt, 2), "volume": round(v_opt, 1), "profit": round((p_opt - unit_cost) * v_opt - fixed_cost)}
    be = fixed_cost / (price - unit_cost) if price > unit_cost else None
    return {"scenarios": rows, "optimal": best, "break_even_volume": round(be, 1) if be is not None else None,
            "note": "" if elasticity < -1 else "poptávka je neelastická (|e| ≤ 1) – zdražení zvyšuje tržby; optimum vzorec nedává, drž se scénářů"}


# ============================================================================================== 443 forecast
def _series(path: Path) -> list[tuple[str, float]]:
    header, rows, _, _ = read_csv(path)
    di = _col(header, "měsíc", "month", "datum", "date", "období", "period")
    vi = _col(header, "tržby", "revenue", "výnos", "obrat", "částka", "amount", "hodnota", "value")
    if di is None or vi is None:
        raise ValueError("Potřebuju sloupce období (měsíc / datum) a hodnota (tržby / částka).")
    agg: dict[str, float] = defaultdict(float)
    for r in rows:
        if max(di, vi) >= len(r):
            continue
        v = _num(r[vi])
        d = _parse_date(r[di]) or (_parse_date(r[di].strip() + "-01") if re.fullmatch(r"\d{4}-\d{1,2}", r[di].strip()) else None)
        if v is None or d is None:
            continue
        agg[f"{d.year}-{d.month:02d}"] += v
    return sorted(agg.items())


def _holt_winters(y: list[float], season: int, horizon: int, a=0.3, b=0.05, g=0.3) -> list[float]:
    if len(y) < 2 * season:
        n = len(y)
        xs = list(range(n))
        mx, my = statistics.fmean(xs), statistics.fmean(y)
        sxx = sum((x - mx) ** 2 for x in xs) or 1
        slope = sum((x - mx) * (v - my) for x, v in zip(xs, y)) / sxx
        return [my + slope * (n + h - mx) for h in range(horizon)]
    level = statistics.fmean(y[:season])
    trend = (statistics.fmean(y[season:2 * season]) - level) / season
    seas = [y[i] - level for i in range(season)]
    for i, v in enumerate(y):
        s = seas[i % season]
        last_level = level
        level = a * (v - s) + (1 - a) * (level + trend)
        trend = b * (level - last_level) + (1 - b) * trend
        seas[i % season] = g * (v - level) + (1 - g) * s
    return [level + (h + 1) * trend + seas[(len(y) + h) % season] for h in range(horizon)]


def forecast(path: Path, months: int = 6, season: int = 12) -> dict:
    """Monthly totals → a forecast (Holt–Winters with yearly season when there are 2+ years, else a trend line), with
    a back-test on the last months so you can see how far to trust it."""
    s = _series(path)
    if len(s) < 6:
        raise ValueError("Na předpověď potřebuju aspoň 6 měsíců dat.")
    y = [v for _, v in s]
    test = min(6, len(y) // 4) or 1
    back = _holt_winters(y[:-test], season, test)
    mape = statistics.fmean(abs(p - a) / abs(a) for p, a in zip(back, y[-test:]) if a) * 100 if any(y[-test:]) else None
    fc = _holt_winters(y, season, months)
    last = dt.date.fromisoformat(s[-1][0] + "-01")
    labels = []
    for h in range(1, months + 1):
        m = last.month - 1 + h
        labels.append(f"{last.year + m // 12}-{m % 12 + 1:02d}")
    method = "Holt–Winters se sezónou" if len(y) >= 2 * season else "lineární trend (na sezónnost je potřeba 2 roky dat)"
    return {"history": s, "forecast": [{"month": m, "value": round(v, 2)} for m, v in zip(labels, fc)], "method": method,
            "backtest_months": test, "mape_pct": round(mape, 1) if mape is not None else None,
            "trust": "dobrá" if mape is not None and mape < 10 else "střední" if mape is not None and mape < 25 else "nízká"}


# ============================================================================================== 444 feedback
POS = ["skvel", "super", "vybor", "spokojen", "rychl", "doporuc", "dekuji", "diky", "perfekt", "paradni", "ochotn", "kvalitn", "libi",
       "uzasn", "bezvadn", "profesional", "great", "excellent", "love", "fast", "perfect", "recommend", "helpful"]
NEG = ["spatn", "pomal", "nespokojen", "problem", "chyb", "nefung", "rozbit", "drah", "zklam", "hroz", "nikdy vic", "reklamac", "nedoruc",
       "zpozd", "nereaguj", "neochot", "bad", "slow", "broken", "terrible", "refund", "never again", "disappoint", "late"]
TOPICS = {"doručení": ["doruc", "dodan", "zasilk", "kuryr", "dopravc", "balik", "delivery", "shipping"],
          "cena": ["cen", "drah", "levn", "sleva", "price", "expensive", "cheap"],
          "kvalita": ["kvalit", "rozbit", "vadn", "nefung", "quality", "broken"],
          "podpora": ["podpor", "zakaznick", "ochot", "reakc", "odpoved", "support", "helpful"],
          "web a aplikace": ["web", "aplikac", "stranka", "objednavk", "kosik", "platb", "app", "checkout"],
          "vrácení a reklamace": ["reklamac", "vracen", "vratil", "refund", "return"]}
NEGATORS = r"\b(ne|nebyl[oa]?|neni|nikdy|not|no)\b"


def sentiment(text: str) -> float:
    t = fold(text)
    words = re.findall(r"\w+", t)
    score = 0.0
    for i, w in enumerate(words):
        sign = 1 if any(w.startswith(p) for p in POS) else -1 if any(w.startswith(n) for n in NEG) else 0
        if sign and any(re.fullmatch(NEGATORS, x) for x in words[max(0, i - 2):i]):
            sign = -sign
        score += sign
    for phrase in ("nikdy vic", "never again"):
        score -= 2 * t.count(phrase)
    return max(-1.0, min(1.0, score / 3))


def feedback(texts: list[str]) -> dict:
    rows = []
    topic_s: dict[str, list[float]] = defaultdict(list)
    for t in texts:
        if not t.strip():
            continue
        s = sentiment(t)
        f = fold(t)
        tops = [k for k, ws in TOPICS.items() if any(w in f for w in ws)]
        for k in tops:
            topic_s[k].append(s)
        rows.append({"text": t.strip()[:300], "score": round(s, 2), "label": "pozitivní" if s > 0.15 else "negativní" if s < -0.15 else "neutrální",
                     "topics": tops})
    if not rows:
        raise ValueError("Žádné texty.")
    topics = sorted(({"topic": k, "mentions": len(v), "avg": round(statistics.fmean(v), 2), "negative": sum(1 for x in v if x < -0.15)}
                     for k, v in topic_s.items()), key=lambda x: (-x["negative"], -x["mentions"]))
    c = Counter(r["label"] for r in rows)
    return {"count": len(rows), "positive": c["pozitivní"], "negative": c["negativní"], "neutral": c["neutrální"],
            "nps_like": round(100 * (c["pozitivní"] - c["negativní"]) / len(rows)), "topics": topics,
            "worst": sorted(rows, key=lambda r: r["score"])[:5], "rows": rows}


def feedback_texts(path: Path) -> list[str]:
    if path.suffix.lower() == ".csv":
        header, rows, _, _ = read_csv(path)
        ti = _col(header, "text", "komentář", "recenze", "zpětná vazba", "feedback", "review", "hodnocení", "zpráva")
        if ti is None:
            ti = max(range(len(header)), key=lambda i: sum(len(r[i]) for r in rows if i < len(r)))
        return [r[ti] for r in rows if ti < len(r)]
    return [x for x in re.split(r"\n\s*\n|\r?\n", path.read_text(encoding="utf-8", errors="replace")) if x.strip()]


# ============================================================================================== 449 profitability
def profitability(path: Path, default_rate: float = 0) -> dict:
    """Time + cost records (project, hours, rate, cost, invoiced) → margin per project, effective hourly rate, and
    which projects lose money."""
    header, rows, _, _ = read_csv(path)
    pi = _col(header, "projekt", "project", "zakázka")
    hi, ri = _col(header, "hodiny", "hours", "čas"), _col(header, "sazba", "rate")
    ci, ii = _col(header, "náklady", "cost", "výdaje"), _col(header, "fakturováno", "invoiced", "příjem", "tržba", "revenue", "cena")
    if pi is None:
        raise ValueError("Potřebuju sloupec projekt.")
    agg: dict[str, dict[str, float]] = defaultdict(lambda: {"hours": 0.0, "labour": 0.0, "cost": 0.0, "revenue": 0.0})

    def val(r, i):
        return (_num(r[i]) or 0.0) if i is not None and i < len(r) else 0.0
    for r in rows:
        if pi >= len(r) or not r[pi].strip():
            continue
        a = agg[r[pi].strip()]
        h = val(r, hi)
        rate = val(r, ri) or default_rate
        a["hours"] += h
        a["labour"] += h * rate
        a["cost"] += val(r, ci)
        a["revenue"] += val(r, ii)
    out = []
    for p, a in agg.items():
        total_cost = a["labour"] + a["cost"]
        profit = a["revenue"] - total_cost
        out.append({"project": p, "hours": round(a["hours"], 1), "revenue": round(a["revenue"]), "cost": round(total_cost), "profit": round(profit),
                    "margin_pct": round(100 * profit / a["revenue"], 1) if a["revenue"] else None,
                    "effective_rate": round((a["revenue"] - a["cost"]) / a["hours"]) if a["hours"] else None})
    out.sort(key=lambda x: x["profit"])
    tot_rev, tot_profit = sum(x["revenue"] for x in out), sum(x["profit"] for x in out)
    return {"projects": out, "losing": [x["project"] for x in out if x["profit"] < 0], "revenue": tot_rev, "profit": tot_profit,
            "margin_pct": round(100 * tot_profit / tot_rev, 1) if tot_rev else None}


def parse_stages(text: str) -> list[tuple[str, float]]:
    """„návštěva 10000, košík 800, objednávka 200“ → [(name, count), …]."""
    out: list[tuple[str, Any]] = []
    for part in re.split(r"[,;\n]|→|->", text):
        m = re.match(r"\s*(.+?)\s*[:=]?\s*(\d[\d\s.,]*)\s*$", part)
        if m:
            out.append((m.group(1).strip(), _num(m.group(2)) or 0.0))
    return out
