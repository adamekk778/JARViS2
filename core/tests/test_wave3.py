"""OMEGA wave 3 (Intelligent File Universe): every function on real files in a temp folder."""
import io
import tarfile
import zipfile
from pathlib import Path

import pytest
from PIL import Image

from jarvis import fileuniverse as fu
from jarvis import tools

tools.load_builtin()


def img(path: Path, color=(200, 30, 30), size=(320, 240), exif_gps=False, **kw):
    im = Image.new("RGB", size, color)
    for x in range(0, size[0], 40):                      # some structure, so the hash has something to see
        for y in range(size[1]):
            im.putpixel((x, y), (255, 255, 255))
    ex = Image.Exif()
    ex[306] = kw.get("date", "2024:07:14 10:20:30")
    ex[271] = "Canon"
    if exif_gps:
        ex[0x8825] = {1: "N", 2: (50.0, 5.0, 0.0)}
    im.save(path, "JPEG", exif=ex.tobytes())
    return path


@pytest.fixture
def work(tmp_path):
    """A clean folder (tmp_path itself also holds JARVIS_HOME)."""
    d = tmp_path / "w"
    d.mkdir()
    return d


@pytest.fixture
def trash_bin(monkeypatch):
    """send2trash stand-in that records instead of touching the real Recycle Bin."""
    import sys
    import types

    got = []
    mod = types.ModuleType("send2trash")

    def send(p):
        got.append(p)
        Path(p).unlink()
    mod.send2trash = send
    monkeypatch.setitem(sys.modules, "send2trash", mod)
    return got


# ---------------------------------------------------------------------------------------------- duplicates
def test_duplicates_by_content_keep_oldest(work):
    (work / "a").mkdir()
    (work / "b").mkdir()
    data = b"x" * 70000 + b"stejne"
    (work / "a" / "original.bin").write_bytes(data)
    import os
    import time

    os.utime(work / "a" / "original.bin", (time.time() - 1000, time.time() - 1000))
    (work / "b" / "kopie (1).bin").write_bytes(data)
    (work / "b" / "jiny.bin").write_bytes(b"x" * 70000 + b"jinaci")     # same size and start, other content
    (work / "b" / "maly.txt").write_bytes(b"ahoj")
    groups = fu.duplicates([work], min_size=1024)
    assert len(groups) == 1 and groups[0]["keep"].endswith("original.bin") and groups[0]["extra"][0].endswith("kopie (1).bin")
    assert fu.wasted(groups) == len(data)


async def test_remove_duplicates_keeps_one_and_uses_the_bin(rt, work, trash_bin):
    for n in ("a.jpg", "b.jpg", "c.jpg"):
        (work / n).write_bytes(b"F" * 5000)
    r = await tools.call("remove_duplicates", {"folder": str(work)}, rt)
    assert r.ok and len(trash_bin) == 2 and len(list(work.iterdir())) == 1


def test_trash_refuses_when_the_kept_copy_changed(work, trash_bin):
    keep, extra = work / "k.bin", work / "e.bin"
    keep.write_bytes(b"1" * 2000)
    extra.write_bytes(b"1" * 2000)
    sha = fu.sha256(keep)
    keep.write_bytes(b"2" * 2000)
    done, failed = fu.trash([str(extra)], {str(extra): (str(keep), sha)})
    assert not done and "nemažu" in failed[0] and extra.exists()


def scene(path: Path, seed: int, size=(320, 240)):
    """A photo-like picture: smooth random blobs (what a difference hash is meant for)."""
    import random

    from PIL import ImageDraw, ImageFilter

    rnd = random.Random(seed)
    im = Image.new("RGB", size, (rnd.randrange(256), rnd.randrange(256), rnd.randrange(256)))
    dr = ImageDraw.Draw(im)
    for _ in range(12):
        x, y = rnd.randrange(size[0]), rnd.randrange(size[1])
        r = rnd.randrange(20, 90)
        dr.ellipse((x - r, y - r, x + r, y + r), fill=(rnd.randrange(256), rnd.randrange(256), rnd.randrange(256)))
    im.filter(ImageFilter.GaussianBlur(6)).save(path, quality=90)
    return path


def test_similar_photos(work):
    a = scene(work / "a.jpg", 1)
    with Image.open(a) as im:
        im.resize((160, 120)).save(work / "a_small.jpg", quality=55)          # smaller + recompressed copy
        im.point(lambda v: min(255, v + 18)).save(work / "a_bright.jpg")     # slightly brighter edit
    scene(work / "other.jpg", 2)
    scene(work / "third.jpg", 3)
    groups = fu.similar_images([work])
    assert [sorted(Path(x).name for x in g) for g in groups] == [["a.jpg", "a_bright.jpg", "a_small.jpg"]]


# ---------------------------------------------------------------------------------------------- rename, sort, names
async def test_rename_by_photo_date_with_preview_and_undo(rt, work):
    img(work / "IMG_2.jpg", date="2024:07:15 09:00:00")
    img(work / "IMG_1.jpg", date="2024:07:14 08:00:00")
    (work / "pozn.txt").write_text("x", encoding="utf-8")
    details = tools.details_for(tools.REGISTRY["rename_files"], {"folder": str(work), "pattern": "Dovolená_{date}_{n}", "only_ext": "jpg"}, rt)
    assert "IMG_1.jpg → Dovolená_2024-07-14_01.jpg" in details
    r = await tools.call("rename_files", {"folder": str(work), "pattern": "Dovolená_{date}_{n}", "only_ext": "jpg"}, rt)
    assert r.ok and sorted(p.name for p in work.iterdir()) == ["Dovolená_2024-07-14_01.jpg", "Dovolená_2024-07-15_02.jpg", "pozn.txt"]
    u = await tools.call("undo_rename", {}, rt)
    assert u.ok and sorted(p.name for p in work.iterdir()) == ["IMG_1.jpg", "IMG_2.jpg", "pozn.txt"]
    assert not (await tools.call("undo_rename", {}, rt)).ok              # nothing left to undo


def test_rename_chains_and_conflicts(work):
    for n in ("a.txt", "b.txt", "c.txt"):
        (work / n).write_text(n, encoding="utf-8")
    (work / "x_1.txt").write_text("cizí", encoding="utf-8")                            # not being renamed → must not be overwritten
    plan = fu.plan_rename(sorted(work.glob("[abc].txt")), "x_{name}", "a", "1")
    fu.apply_rename(plan)
    assert (work / "x_1.txt").read_text(encoding="utf-8") == "cizí" and (work / "x_1 (2).txt").read_text(encoding="utf-8") == "a.txt"
    # swap: a → b, b → a works thanks to the two-phase rename
    (work / "s1").write_text("1", encoding="utf-8")
    (work / "s2").write_text("2", encoding="utf-8")
    fu.apply_rename([{"src": str(work / "s1"), "dst": str(work / "s2")}, {"src": str(work / "s2"), "dst": str(work / "s1")}])
    assert (work / "s1").read_text(encoding="utf-8") == "2" and (work / "s2").read_text(encoding="utf-8") == "1"


def _docx(path: Path, text: str):
    import docx

    d = docx.Document()
    for line in text.splitlines():
        d.add_paragraph(line)
    d.core_properties.author = "Jan Novák"
    d.save(str(path))
    return path


async def test_sort_documents_by_content_and_undo(rt, work):
    _docx(work / "dokument1.docx", "Faktura – daňový doklad č. 2026001\nVariabilní symbol 2026001")
    (work / "scan.txt").write_text("Výpis z účtu za září, zůstatek 12 000 Kč", encoding="utf-8")
    img(work / "foto.jpg")
    (work / "setup.exe").write_bytes(b"MZ")
    d = tools.details_for(tools.REGISTRY["sort_documents"], {"folder": str(work)}, rt)
    assert "dokument1.docx → Faktury/ (podle obsahu)" in d
    r = await tools.call("sort_documents", {"folder": str(work)}, rt)
    assert r.ok and (work / "Faktury" / "dokument1.docx").exists() and (work / "Banka a výpisy" / "scan.txt").exists()
    assert (work / "Fotky" / "foto.jpg").exists() and (work / "Instalačky" / "setup.exe").exists()
    await tools.call("undo_rename", {}, rt)
    assert (work / "dokument1.docx").exists() and (work / "setup.exe").exists()


def test_smart_name_from_content(work):
    p = _docx(work / "scan0001.docx", "FAKTURA\nDodavatel: Alza.cz a.s.\nDatum vystavení: 3. 10. 2026")
    assert fu.smart_name(p) == "Faktura_2026-10-03_Alza.cz.docx"
    q = work / "poznamky.txt"
    q.write_text("nic zajímavého", encoding="utf-8")
    assert fu.smart_name(q) == "poznamky.txt"


# ---------------------------------------------------------------------------------------------- archives
def test_archive_list_and_zip_slip_protection(work):
    z = work / "evil.zip"
    with zipfile.ZipFile(z, "w") as f:
        f.writestr("ok/readme.txt", "ahoj")
        f.writestr("../../escape.txt", "zlo")
    info = fu.archive_list(z)
    assert info["type"] == "ZIP" and info["files"] == 2 and not info["encrypted"]
    r = fu.safe_extract(z, work / "out")
    assert r["files"] == 1 and r["skipped"] == ["../../escape.txt"] and not (work / "escape.txt").exists()
    with pytest.raises(ValueError):
        fu.safe_extract(z, work / "out")                  # not empty → never mixes into existing files


def test_tar_and_make_zip_roundtrip(work):
    src = work / "projekt"
    (src / "sub").mkdir(parents=True)
    (src / "a.txt").write_text("A", encoding="utf-8")
    (src / "sub" / "b.txt").write_text("B", encoding="utf-8")
    r = fu.make_zip([src], work / "p.zip")
    assert r["files"] == 2
    out = fu.safe_extract(work / "p.zip", work / "x")
    assert (work / "x" / "projekt" / "sub" / "b.txt").read_text(encoding="utf-8") == "B" and out["files"] == 2
    t = work / "a.tar.gz"
    with tarfile.open(t, "w:gz") as tf:
        info = tarfile.TarInfo("data/c.txt")
        info.size = 1
        tf.addfile(info, io.BytesIO(b"C"))
    assert fu.archive_list(t)["files"] == 1
    assert fu.safe_extract(t, work / "t")["files"] == 1 and (work / "t" / "data" / "c.txt").read_text(encoding="utf-8") == "C"
    assert "error" in fu.archive_list(work / "projekt" / "a.txt")


async def test_archive_tools(rt, work):
    (work / "f.txt").write_text("x", encoding="utf-8")
    r = await tools.call("create_zip", {"paths": [str(work / "f.txt")]}, rt)
    assert r.ok and (work / "f.zip").exists()
    c = await tools.call("archive_contents", {"path": str(work / "f.zip")}, rt)
    assert c.ok and "f.txt" in c.summary
    e = await tools.call("extract_archive", {"path": str(work / "f.zip")}, rt)
    assert e.ok and (work / "f (2)" / "f.txt").exists() or (work / "f" / "f.txt").exists()


# ---------------------------------------------------------------------------------------------- conversion
def test_convert_images_pdf_word_excel(work):
    p = work / "logo.png"
    Image.new("RGBA", (800, 400), (0, 0, 255, 128)).save(p)
    j = fu.convert(p, "jpg", max_side=200)
    with Image.open(j) as im:
        assert im.format == "JPEG" and max(im.size) == 200
    assert fu.convert(p, "pdf").read_bytes().startswith(b"%PDF")
    d = _docx(work / "a.docx", "Nadpis\nřádek dva")
    assert "řádek dva" in fu.convert(d, "txt").read_text(encoding="utf-8")
    csvp = work / "t.csv"
    csvp.write_text("jméno;částka\nJana;1 200,50\nPetr;300\n", encoding="utf-8")
    x = fu.convert(csvp, "xlsx")
    import openpyxl

    ws = openpyxl.load_workbook(x).active
    assert ws["B2"].value == 1200.5 and ws["B3"].value == 300
    back = fu.convert(x, "csv").read_text(encoding="utf-8-sig")
    assert "Jana;1200.5" in back
    md = work / "n.md"
    md.write_text("# Titul\n- bod\n", encoding="utf-8")
    import docx

    assert [par.text for par in docx.Document(str(fu.convert(md, "docx"))).paragraphs] == ["Titul", "bod"]
    with pytest.raises(ValueError):
        fu.convert(csvp, "mp3")
    assert p.exists()                                       # originals stay


def test_pdf_to_text(work):
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(100, 100)
    with open(work / "blank.pdf", "wb") as f:
        w.write(f)
    with pytest.raises(ValueError, match="naskenované"):
        fu.convert(work / "blank.pdf", "txt")


# ---------------------------------------------------------------------------------------------- metadata, versions
async def test_metadata_read_and_clean(rt, work):
    p = img(work / "dovolena.jpg", exif_gps=True)
    md = await tools.call("file_metadata", {"path": str(p)}, rt)
    assert "GPS" in md.summary and "Canon" in md.summary
    r = await tools.call("clean_metadata", {"path": str(p)}, rt)
    clean = work / "dovolena_bez_metadat.jpg"
    assert r.ok and clean.exists() and fu.read_metadata(clean) == {} and "GPS" in fu.read_metadata(p)
    d = _docx(work / "cv.docx", "text")
    out = fu.clean_metadata(d)
    assert "author" not in fu.read_metadata(out) and fu.read_metadata(d)["author"] == "Jan Novák"
    assert tools.is_dangerous(tools.REGISTRY["clean_metadata"], {"overwrite": True}) and not tools.is_dangerous(tools.REGISTRY["clean_metadata"], {})


def test_metadata_overwrite_keeps_a_version(work):
    p = img(work / "x.jpg", exif_gps=True)
    before = fu.sha256(p)
    fu.clean_metadata(p, overwrite=True)
    assert fu.read_metadata(p) == {} and fu.list_versions(p)[0]["sha256"] == before


async def test_versions_save_list_restore(rt, work):
    p = work / "plan.txt"
    p.write_text("v1", encoding="utf-8")
    await tools.call("save_file_version", {"path": str(p)}, rt)
    assert (await tools.call("save_file_version", {"path": str(p)}, rt)).summary.startswith("Soubor se od poslední verze nezměnil")
    p.write_text("v2 – pokažené", encoding="utf-8")
    r = await tools.call("restore_file_version", {"path": str(p), "index": 0}, rt)
    assert r.ok and p.read_text(encoding="utf-8") == "v1" and "ověřeno" in r.summary
    assert fu.list_versions(p)[0]["sha256"] == fu.sha256(p) or len(fu.list_versions(p)) == 2   # the broken state was saved too
    assert any(v["note"] == "před obnovením starší verze" for v in fu.list_versions(p))


# ---------------------------------------------------------------------------------------------- integrity, lost files
async def test_folder_fingerprint_and_check(rt, work):
    f = work / "foto"
    f.mkdir()
    (f / "a.jpg").write_bytes(b"A")
    (f / "b.jpg").write_bytes(b"B")
    assert not (await tools.call("folder_check", {"folder": str(f)}, rt)).ok
    await tools.call("folder_fingerprint", {"folder": str(f)}, rt)
    assert "beze změny" in (await tools.call("folder_check", {"folder": str(f)}, rt)).summary
    (f / "a.jpg").write_bytes(b"A2")
    (f / "b.jpg").unlink()
    (f / "c.jpg").write_bytes(b"C")
    r = await tools.call("folder_check", {"folder": str(f)}, rt)
    assert not r.ok and r.data["changed"] == ["a.jpg"] and r.data["missing"] == ["b.jpg"] and r.data["added"] == ["c.jpg"]


def test_recent_and_content_search(work):
    import os
    import time

    (work / "stary.txt").write_text("nic", encoding="utf-8")
    os.utime(work / "stary.txt", (time.time() - 10 * 86400,) * 2)
    _docx(work / "smlouva.docx", "Nájemní smlouva na byt v Brně, nájemné 15 000 Kč")
    rec = fu.recent_files([work], hours=1)
    assert [Path(r["path"]).name for r in rec] == ["smlouva.docx"]
    hits = fu.search_content([work], "nájemné Brně")
    assert len(hits) == 1 and hits[0]["path"].endswith("smlouva.docx") and "nájemné" in hits[0]["snippet"]


# ---------------------------------------------------------------------------------------------- backup
async def test_backup_incremental_and_verified(rt, work):
    src, dst = work / "Dokumenty", work / "Disk_E"
    (src / "sub").mkdir(parents=True)
    (src / "a.txt").write_text("A", encoding="utf-8")
    (src / "sub" / "b.txt").write_text("B", encoding="utf-8")
    r1 = await tools.call("backup_folder", {"source": str(src), "destination": str(dst)}, rt)
    assert r1.ok and r1.data["copied"] == 2 and (dst / "Dokumenty" / "sub" / "b.txt").read_text(encoding="utf-8") == "B"
    import os
    import time

    (src / "a.txt").write_text("A2", encoding="utf-8")
    os.utime(src / "a.txt", (time.time() + 5,) * 2)
    r2 = await tools.call("backup_folder", {"source": str(src), "destination": str(dst)}, rt)
    assert r2.data["copied"] == 1 and r2.data["unchanged"] == 1 and (dst / "Dokumenty" / "a.txt").read_text(encoding="utf-8") == "A2"
    bad = await tools.call("backup_folder", {"source": str(src), "destination": str(src / "zaloha")}, rt)
    assert not bad.ok


def test_backup_plan_orders_by_value(work):
    for name in ("Downloads", "Documents", "Pictures"):
        (work / name).mkdir()
        (work / name / "f.bin").write_bytes(b"x" * 100)
    (work / "Downloads" / "setup.exe").write_bytes(b"x" * 10000)
    plan = fu.backup_plan(work)
    assert [Path(p["folder"]).name for p in plan] == ["Documents", "Pictures", "Downloads"] and plan[-1]["note"]


def test_read_only_mode_allows_the_finders():
    from jarvis.permissions import READ_ONLY_OK

    for n in ("find_duplicates", "find_similar_photos", "archive_contents", "folder_check", "find_in_files"):
        assert n in READ_ONLY_OK
    for n in ("remove_duplicates", "rename_files", "sort_documents", "smart_rename", "undo_rename", "backup_folder", "restore_file_version"):
        assert tools.REGISTRY[n].danger, n


# ---------------------------------------------------------------------------------------------- the page (WebSocket)
def test_files_page_reads_directly_and_asks_before_changes(rt, work):
    from fastapi.testclient import TestClient

    from jarvis.server.app import create_app

    for n in ("a.bin", "b.bin"):
        (work / n).write_bytes(b"D" * 3000)
    with TestClient(create_app(rt, manage_runtime=False), client=("127.0.0.1", 5000)) as c, c.websocket_connect("/ws") as ws:
        ws.send_json({"type": "files_run", "tool": "find_duplicates", "args": {"folder": str(work)}})
        m = next(x for x in iter(ws.receive_json, None) if x["type"] == "files_result")
        assert m["ok"] and m["data"][0]["extra"]
        ws.send_json({"type": "files_run", "tool": "run_shell", "args": {"command": "dir"}})
        assert "nejde" in next(x for x in iter(ws.receive_json, None) if x["type"] == "error")["message"]
        ws.send_json({"type": "files_run", "tool": "remove_duplicates", "args": {"folder": str(work)}})
        req = next(x for x in iter(ws.receive_json, None) if x["type"] == "confirm_request")
        assert any("Do koše: 1 souborů" in d for d in req["details"])
        ws.send_json({"type": "confirm_response", "id": req["id"], "approved": False})
        m = next(x for x in iter(ws.receive_json, None) if x["type"] == "files_result")
    assert (work / "a.bin").exists() and (work / "b.bin").exists()
