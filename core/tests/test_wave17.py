"""OMEGA wave 17 (business): RFM segments, lead scoring, funnel, a quote with VAT as a real .docx (read back), pricing
scenarios against the closed-form optimum, a seasonal revenue forecast, feedback sentiment/topics, profitability."""
import csv
import datetime as dt
import math

import pytest

from jarvis import business, tools

tools.load_builtin()


def write(path, rows, delim=";"):
    with open(path, "w", encoding="utf-8", newline="") as f:
        csv.writer(f, delimiter=delim).writerows(rows)
    return path


async def test_customer_segments(rt, tmp_path):
    today = dt.date(2026, 10, 1)
    rows = [["zákazník", "datum", "částka"]]
    for i in range(6):                                                  # champions: recent, often, a lot
        rows += [[f"vip{i}", (today - dt.timedelta(days=d)).isoformat(), "2000"] for d in (2, 20, 40, 60, 80, 100)]
    for i in range(6):                                                  # lost: one small order long ago
        rows.append([f"old{i}", (today - dt.timedelta(days=400 + i)).isoformat(), "150"])
    for i in range(6):                                                  # at risk: used to buy often, not lately
        rows += [[f"risk{i}", (today - dt.timedelta(days=d)).isoformat(), "800"] for d in (250, 280, 300, 320)]
    for i in range(6):                                                  # new: one recent order
        rows.append([f"new{i}", (today - dt.timedelta(days=3)).isoformat(), "500"])
    rows.append(["", "2026-01-01", "x"])
    r = await tools.call("customer_segments", {"file": str(write(tmp_path / "o.csv", rows))}, rt)
    seg = {c["customer"]: c["segment"] for c in r.data["customers"]}
    assert seg["vip0"] == "šampioni" and seg["old0"] == "ztracení" and seg["risk0"] == "ohrožení" and seg["new0"] == "noví"
    assert r.data["skipped_rows"] == 1 and r.data["segments"][0]["segment"] == "šampioni"
    bad = write(tmp_path / "b.csv", [["a", "b"], ["1", "2"]])
    assert not (await tools.call("customer_segments", {"file": str(bad)}, rt)).ok


async def test_lead_scoring(rt, tmp_path):
    f = write(tmp_path / "leads.csv", [["firma", "rozpočet", "role", "termín", "zaměstnanci", "návštěvy"],
                                       ["Alfa s.r.o.", "150000", "CEO", "do 2 měsíců", "120", "12"],
                                       ["Beta", "20000", "asistentka", "příští rok", "5", "1"],
                                       ["Gama", "", "", "", "", ""]])
    r = await tools.call("score_leads", {"file": str(f)}, rt)
    leads = r.data["leads"]
    assert [x["lead"] for x in leads] == ["Alfa s.r.o.", "Beta", "Gama"] and leads[0]["grade"] == "horký" and leads[2]["score"] == 0
    assert set(leads[0]["why"]) >= {"rozhoduje", "brzy", "aktivní"} and r.data["hot"] == 1


async def test_funnel(rt, tmp_path):
    r = await tools.call("sales_funnel", {"stages": "návštěva 10000, košík 800, objednávka 200, platba 190"}, rt)
    assert r.data["overall_pct"] == 1.9 and r.data["leak"]["from"] == "návštěva" and r.data["leak"]["rate_pct"] == 8.0
    assert r.data["if_leak_plus10"] == pytest.approx(19.0)
    assert not (await tools.call("sales_funnel", {"stages": "a 10, b 20"}, rt)).ok
    ev = write(tmp_path / "ev.csv", [["user", "stage"], *[[f"u{i}", "visit"] for i in range(10)], *[[f"u{i}", "cart"] for i in range(4)],
                                     ["u0", "buy"], ["u9", "buy"]])
    r = await tools.call("sales_funnel", {"stages": "visit, cart, buy", "file": str(ev)}, rt)
    assert [s for _, s in [(x["to"], x["rate_pct"]) for x in r.data["steps"]]] == [40.0, 25.0]   # u9 bought without a cart → not counted


async def test_price_quote_docx(rt, tmp_path):
    import docx

    items = [{"name": "Tvorba webu", "qty": 1, "price": 25000}, {"name": "Hosting", "qty": 12, "unit": "měs.", "price": 150},
             {"name": "Kniha", "qty": 2, "price": 400, "vat": 12}]
    r = await tools.call("price_quote", {"client": "Pekárna Novák", "items": items, "supplier": "Adam Bartok", "discount_pct": 10,
                                         "target": str(tmp_path / "nabidka.docx")}, rt)
    assert r.ok and r.data["base"] == pytest.approx(0.9 * (25000 + 1800 + 800))
    assert r.data["by_rate"][21.0]["vat"] == pytest.approx(0.9 * 26800 * 0.21) and r.data["by_rate"][12.0]["vat"] == pytest.approx(86.4)
    text = "\n".join(p.text for p in docx.Document(str(tmp_path / "nabidka.docx")).paragraphs)
    assert "Odběratel: Pekárna Novák" in text and "Celkem k úhradě: 29 991,60 Kč" in text and "slevu 10 %" in text
    assert tools.REGISTRY["price_quote"].perm == "files_write"
    r = await tools.call("price_quote", {"client": "X", "items": [{"name": "a", "price": 100}], "vat_payer": False,
                                         "target": str(tmp_path / "n2.docx")}, rt)
    assert r.data["total"] == 100 and "není plátcem DPH" in "\n".join(p.text for p in docx.Document(str(tmp_path / "n2.docx")).paragraphs)
    assert not (await tools.call("price_quote", {"client": "X", "items": [{"name": "a", "price": 1}], "target": str(tmp_path / "n2.docx")}, rt)).ok


def test_pricing_matches_closed_form():
    r = business.price_scenarios(100, 1000, 60, fixed_cost=10000, elasticity=-2)
    assert r["optimal"]["price"] == pytest.approx(120)               # p* = c·e/(1+e) = 60·2
    base = next(s for s in r["scenarios"] if s["change_pct"] == 0)
    assert base["profit"] == 30000 and r["break_even_volume"] == 250
    up10 = next(s for s in r["scenarios"] if s["change_pct"] == 10)
    assert up10["volume"] == pytest.approx(1000 * 1.1 ** -2, rel=1e-3)
    assert business.price_scenarios(100, 1000, 60, elasticity=-0.5)["optimal"] is None


async def test_revenue_forecast_seasonal(rt, tmp_path):
    rows = [["měsíc", "tržby"]]
    for y in (2023, 2024, 2025):
        for m in range(1, 13):
            rows.append([f"{y}-{m:02d}", str(round(100000 + (y - 2023) * 12000 + m * 1000 + (40000 if m == 12 else 0)))])
    r = await tools.call("revenue_forecast", {"file": str(write(tmp_path / "t.csv", rows)), "months": 12}, rt)
    fc = {f["month"]: f["value"] for f in r.data["forecast"]}
    assert r.data["method"].startswith("Holt") and len(fc) == 12 and list(fc)[0] == "2026-01"
    assert fc["2026-12"] > fc["2026-11"] + 25000 and fc["2026-06"] > 130000          # December peak and growth kept
    assert r.data["mape_pct"] < 10 and r.data["trust"] == "dobrá"
    short = write(tmp_path / "s.csv", [["datum", "částka"], ["1.1.2026", "10"], ["1.2.2026", "20"], ["1.3.2026", "30"], ["1.4.2026", "40"],
                                       ["1.5.2026", "50"], ["1.6.2026", "60"], ["15.6.2026", "5"]])
    r = await tools.call("revenue_forecast", {"file": str(short), "months": 2}, rt)
    assert "lineární trend" in r.data["method"] and r.data["forecast"][0]["value"] == pytest.approx(73.7, abs=3)


async def test_feedback(rt, tmp_path):
    texts = ["Skvělá obsluha, doporučuji!", "Doručení trvalo 3 týdny, zásilka přišla rozbitá. Nikdy víc.", "Cena je moc drahá.",
             "Podpora nereaguje na e-maily.", "Kvalita super, rychlé dodání, děkuji", "Nebylo to špatné.", "Objednávka v pořádku."]
    r = await tools.call("analyze_feedback", {"texts": texts}, rt)
    labels = [x["label"] for x in r.data["rows"]]
    assert labels[:5] == ["pozitivní", "negativní", "negativní", "negativní", "pozitivní"] and labels[5] == "pozitivní"
    tops = {t["topic"]: t for t in r.data["topics"]}
    assert tops["doručení"]["mentions"] == 2 and tops["doručení"]["negative"] == 1 and "podpora" in tops
    assert r.data["worst"][0]["text"].startswith("Doručení")
    f = write(tmp_path / "rec.csv", [["hvězdy", "recenze"], ["5", "Super rychlé"], ["1", "Špatná kvalita"]])
    assert (await tools.call("analyze_feedback", {"file": str(f)}, rt)).data["count"] == 2
    assert not (await tools.call("analyze_feedback", {"texts": []}, rt)).ok


async def test_profitability(rt, tmp_path):
    f = write(tmp_path / "vykaz.csv", [["projekt", "hodiny", "sazba", "náklady", "fakturováno"],
                                       ["Web Novák", "40", "500", "1000", "35000"], ["Web Novák", "10", "500", "", ""],
                                       ["E-shop Beta", "120", "500", "5000", "50000"], ["Logo", "6", "", "", "6000"]])
    r = await tools.call("project_profitability", {"file": str(f), "hourly_rate": 400}, rt)
    by = {p["project"]: p for p in r.data["projects"]}
    assert by["Web Novák"]["profit"] == 35000 - 25000 - 1000 and by["Web Novák"]["effective_rate"] == round(34000 / 50)
    assert by["E-shop Beta"]["profit"] == -15000 and r.data["losing"] == ["E-shop Beta"] and by["Logo"]["profit"] == 6000 - 2400
    assert math.isclose(r.data["profit"], 9000 - 15000 + 3600)
