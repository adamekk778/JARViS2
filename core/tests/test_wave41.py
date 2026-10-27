"""Wave 41: next-step prediction, contradictions, improving code until tests pass, repeated work, priority negotiation,
handing data between apps, app capabilities, the semantic desktop index, input mode and answer sources."""
import datetime as dt
import sys
import time

from openpyxl import load_workbook

from conftest import FakeRouter
from jarvis import intelplus, proactive, tools

tools.load_builtin()


def rows(seqs, start=1_000_000.0):
    out, t = [], start
    for s in seqs:
        for tool in s:
            out.append({"ts": t, "tool": tool})
            t += 60
        t += 3 * 3600
    return out


DAY = [["mail_check", "read_pdf", "invoice_add", "send_to_phone"]] * 4 + [["mail_check", "weather"]]


def test_sessions_and_prediction():
    seqs = intelplus.sessions(rows(DAY + [["update_plan", "mail_check"]]))
    assert len(seqs) == 6 and seqs[-1] == ["mail_check"]                    # bookkeeping tools are ignored
    model = intelplus.markov(seqs)
    best = intelplus.predict(model, ["mail_check"], 3)[0]
    assert best["steps"] == ["read_pdf", "invoice_add", "send_to_phone"] and 0.75 < best["p"] < 0.85
    assert intelplus.predict(model, ["unknown"], 2) == []


def test_repeated_patterns_prefers_longest():
    pats = intelplus.repeated_patterns(intelplus.sessions(rows(DAY)))
    assert pats[0]["steps"] == ["mail_check", "read_pdf", "invoice_add", "send_to_phone"] and pats[0]["sessions"] == 4
    assert not any(p["steps"] == ["read_pdf", "invoice_add"] for p in pats)


def test_app_habits():
    mon = dt.date(2026, 10, 19)
    data = [{"day": (mon - dt.timedelta(days=7 * i)).isoformat(), "hour": 9, "app": "Outlook", "seconds": 1800} for i in range(3)]
    data.append({"day": mon.isoformat(), "hour": 9, "app": "Excel", "seconds": 600})
    h = intelplus.app_habits(data, 0, 9)
    assert h[0] == {"app": "Outlook", "minutes_avg": 30.0} and h[1]["app"] == "Excel"
    assert intelplus.app_habits(data, 1, 9) == []


async def test_prediction_and_opportunity_tools(rt):
    now = time.time()
    rt.tasks  # make sure the schema exists
    for r in rows(DAY, now - 5 * 86400):
        rt.db.exec("INSERT INTO task_actions (task_id, ts, tool, args, ok, summary, ms) VALUES ('t', ?, ?, '{}', 1, '', 1)", (r["ts"], r["tool"]))
    rt.db.exec("INSERT INTO task_actions (task_id, ts, tool, args, ok, summary, ms) VALUES ('t', ?, 'mail_check', '{}', 1, '', 1)", (now - 60,))
    r = await tools.call("predict_next_steps", {}, rt)
    assert r.ok and r.data["paths"][0]["steps"][0] == "read_pdf"
    r = await tools.call("automation_opportunities", {}, rt)
    assert r.ok and r.data[0]["sessions"] == 4
    sug = proactive.next_actions(rt, now=now)
    assert any("dovednost" in s["text"] for s in sug)


def test_contradictions():
    a = "Smlouva o dílo. Splatnost faktury je 30 dní od doručení. Záruka činí 24 měsíců. Dodavatel smí použít subdodavatele."
    b = "Faktura má splatnost 14 dní. Záruka 24 měsíců. Dodavatel nesmí použít subdodavatele."
    found = intelplus.contradictions([{"name": "smlouva.pdf", "text": a}, {"name": "dodatek.docx", "text": b}])
    kinds = sorted(c["kind"] for c in found)
    assert kinds == ["jiná hodnota", "popření"]
    val = next(c for c in found if c["kind"] == "jiná hodnota")
    assert {val["a"]["value"], val["b"]["value"]} == {"30 dní", "14 dní"} and val["b"]["doc"] == "dodatek.docx"
    assert intelplus.contradictions([{"name": "x", "text": "Cena je 100 Kč. Hmotnost je 100 kg."}]) == []
    assert intelplus.contradictions([{"name": "x", "text": "Lhůta je 5 dní. Později: lhůta je 7 dní."}])


async def test_contradictions_tool(rt, tmp_path):
    f = tmp_path / "zadani.txt"
    f.write_text("Rozpočet projektu je 500 000 Kč.\nV závěru: rozpočet projektu je 450 000 Kč.", encoding="utf-8")
    r = await tools.call("find_contradictions", {"files": [str(f)]}, rt)
    assert r.ok and "1 rozporů" in r.summary and "450 000" in r.summary


async def test_improve_loop(tmp_path):
    proj = tmp_path / "p"
    proj.mkdir()
    (proj / "calc.py").write_text("def add(a, b):\n    return a - b\n")
    (proj / "test_calc.py").write_text("from calc import add\nassert add(2, 3) == 5\n")
    seen = []

    async def propose(current, output, n):
        seen.append(output)
        return {"calc.py": current["calc.py"].replace("a - b", "a + b"), "test_calc.py": "assert True\n"}
    r = await intelplus.improve_loop(proj, f"{sys.executable} test_calc.py", ["calc.py"], propose)
    assert r["ok"] and r["rounds"] == 1 and "AssertionError" in seen[0]
    assert "-    return a - b" in r["diff"] and "+    return a + b" in r["diff"]
    assert (proj / "calc.py").read_text().endswith("a - b\n")              # the original stays untouched
    assert (proj / "test_calc.py").read_text().startswith("from calc")    # tests are never changed

    async def nothing(current, output, n):
        return {}
    bad = await intelplus.improve_loop(proj, f"{sys.executable} test_calc.py", ["calc.py"], nothing, 3)
    assert not bad["ok"] and bad["diff"] == "" and bad["log"][-1]["output"] == "AI nenavrhla žádnou změnu."


async def test_improve_tool(rt, tmp_path):
    proj = tmp_path / "proj"
    proj.mkdir()
    (proj / "m.py").write_text("X = 1\n")
    (proj / "t.py").write_text("import m\nassert m.X == 2\n")
    rt.cfg.tools.allowed_folders = [str(tmp_path)]
    rt.router = FakeRouter(['{"files": {"m.py": "X = 2\\n"}}'])
    r = await tools.call("improve_until_tests_pass", {"project": str(proj), "test_cmd": f"{sys.executable} t.py", "files": ["m.py"]}, rt)
    assert r.ok and "PROŠLY" in r.summary and (proj / "m.py").read_text() == "X = 1\n"
    assert tools.REGISTRY["improve_until_tests_pass"].danger


def test_negotiate():
    start = dt.date(2026, 10, 26)                                     # Monday
    tasks = [{"name": "Nabídka", "hours": 10, "due": "2026-10-27", "priority": 1},
             {"name": "Web", "hours": 8, "due": "2026-10-28", "priority": 3},
             {"name": "Report", "hours": 6, "due": "2026-10-28", "priority": 2}]
    plan = intelplus.schedule(tasks, 6, start)
    assert [p["name"] for p in plan] == ["Nabídka", "Report", "Web"] and plan[-1]["late_days"] == 1
    r = intelplus.negotiate(tasks, 6, start)
    assert r["late"] == ["Web"] and r["ok"] and r["postpone"][0]["name"] == "Web" and r["postpone"][0]["new_finish"] == "2026-10-29"
    assert r["extra_hours"] == 6.0
    ok = intelplus.negotiate(tasks[:1], 6, start)
    assert ok["ok"] and not ok["postpone"]
    weekend = intelplus.schedule([{"name": "x", "hours": 12, "due": "2026-11-02"}], 6, dt.date(2026, 10, 30))
    assert weekend[0]["finish"] == "2026-11-02"


def test_text_conversions():
    assert intelplus.table_from_text("Jméno\tCena\nKafe\t45\nČaj\t30") == [["Jméno", "Cena"], ["Kafe", "45"], ["Čaj", "30"]]
    assert intelplus.table_from_text("a;b\n1;2\n") == [["a", "b"], ["1", "2"]]
    assert intelplus.table_from_text("jen text") == []
    now = dt.datetime(2026, 10, 27, 10, 0).timestamp()
    ev = intelplus.event_from_text("Porada s Novákem 3. 11. v 14:30 v Brně", now)
    assert dt.datetime.fromtimestamp(ev["start"]) == dt.datetime(2026, 11, 3, 14, 30) and ev["location"] == "Brně" and "Porada" in ev["title"]
    assert intelplus.event_from_text("zítra v 9 zubař", now)["start"] == dt.datetime(2026, 10, 28, 9, 0).timestamp()
    assert intelplus.event_from_text("bez data", now) is None
    assert intelplus.items_from_text("- mléko\n- chleba\n3) máslo") == ["mléko", "chleba", "máslo"]
    assert intelplus.items_from_text("mléko, vejce, sýr") == ["mléko", "vejce", "sýr"]
    eml = intelplus.eml_draft("a@b.cz", "Ahoj", "text").decode()
    assert "X-Unsent: 1" in eml and "To: a@b.cz" in eml


async def test_cross_app_tool(rt):
    rt.platform.clipboard = "Položka\tKs\nŠrouby\t120\nMatky\t80"
    r = await tools.call("cross_app", {"target": "tabulka"}, rt)
    assert r.ok
    ws = load_workbook(r.attachments[0]).active
    assert ws["B2"].value == 120 and ws["A3"].value == "Matky"
    rt.platform.clipboard = "Schůzka s účetní zítra v 10:00"
    r = await tools.call("cross_app", {"target": "kalendář"}, rt)
    assert r.ok and rt.db.one("SELECT title FROM cal_events")["title"].startswith("Schůzka s účetní")
    rt.platform.clipboard = "- koupit toner\n- zavolat servis"
    r = await tools.call("cross_app", {"target": "seznam", "list_name": "práce"}, rt)
    assert r.ok and r.data["items"] == ["koupit toner", "zavolat servis"]
    r = await tools.call("cross_app", {"target": "e-mail", "to": "x@y.cz"}, rt)
    assert r.ok and "odešleš ho ty" in r.summary and r.attachments[0].endswith(".eml")
    assert not (await tools.call("cross_app", {"target": "fax"}, rt)).ok


HELP = """Usage: tool [options] <command>

Options:
  -h, --help            show help
  -o, --output <file>   where to write
  --quiet               no output

Commands:
  build      build the project
  serve      start a server
"""


def test_parse_help_and_associations():
    r = intelplus.parse_help(HELP)
    assert [o["flags"] for o in r["options"]] == [["-h", "--help"], ["-o", "--output"], ["--quiet"]]
    assert r["options"][1]["arg"] == "file" and [c["name"] for c in r["commands"]] == ["build", "serve"]
    reg = [{"kind": "ext", "name": ".pdf", "command": '"C:\\Program Files\\Acrobat\\Acrobat.exe" "%1"'},
           {"kind": "protocol", "name": "acrobat", "command": "C:\\Program Files\\Acrobat\\Acrobat.exe /u %1"},
           {"kind": "ext", "name": ".txt", "command": "notepad.exe %1"}]
    a = intelplus.app_associations("Acrobat.exe", lambda: reg)
    assert a == {"extensions": [".pdf"], "protocols": ["acrobat"]}


async def test_app_capabilities_tool(rt, tmp_path):
    exe = tmp_path / "mytool"
    exe.write_text(f"#!{sys.executable}\nprint('''{HELP}''')\n")
    exe.chmod(0o755)
    r = await tools.call("app_capabilities", {"app": str(exe)}, rt)
    assert r.ok and "build" in r.summary and r.data["options"]


def test_entities():
    e = intelplus.entities("Dobrý den, Jan Novák (jan.novak@firma.cz, +420 777 123 456) z projektu Atlas, IČO: 12345678.", ["Atlas"])
    assert {("email", "jan.novak@firma.cz"), ("telefon", "777123456"), ("jméno", "jan novak"), ("ičo", "12345678"), ("projekt", "atlas")} <= e
    assert not any(v.startswith("dobry") for _, v in e)


async def test_semantic_index(rt):
    rt.calendar.add("Schůzka Jan Novák – rozpočet", dt.datetime(2026, 11, 3, 10, 0))
    rt.knowledge.add_fact("Jan Novák je účetní, e-mail jan.novak@firma.cz")
    rt.db.exec("CREATE TABLE IF NOT EXISTS lists (id INTEGER PRIMARY KEY AUTOINCREMENT, list TEXT NOT NULL, item TEXT NOT NULL, ts REAL NOT NULL)")
    rt.db.exec("INSERT INTO lists (list, item, ts) VALUES ('práce', 'poslat Janu Novákovi výpis', 1)")
    st = intelplus.build_index(rt)
    assert st["links"] >= 3
    r = intelplus.related(rt, "Novák")
    assert set(r["items"]) >= {"kalendář", "paměť", "seznam"} and "jan.novak@firma.cz" in r["entities"]
    assert any(x["entity"] == "jan.novak@firma.cz" for x in intelplus.related(rt, "jan novak")["linked"])
    t = await tools.call("desktop_links", {"query": "novak"}, rt)
    assert t.ok and "kalendář" in t.summary
    assert intelplus.related(rt, "nikdo")["items"] == {}


def test_input_profile():
    assert intelplus.input_profile({"coarse": True, "hover": False})["mode"] == "touch"
    assert intelplus.input_profile({"last_input": "mouse"})["mode"] == "mouse"
    assert intelplus.input_profile({"last_input": "voice", "fullscreen": True})["mode"] == "voice"
    assert intelplus.input_profile({"preference": "touch"})["target_px"] == 48


def test_sources_from():
    s = intelplus.sources_from("search_documents", {"query": "x"}, [{"file": "C:/docs/smlouva.pdf", "loc": "s. 3", "text": "…"}])
    assert s == [{"kind": "soubor", "ref": "C:/docs/smlouva.pdf (s. 3)", "label": "smlouva.pdf, s. 3", "tool": "search_documents"}]
    w = intelplus.sources_from("web_search", {"query": "x"}, {"results": [{"url": "https://a.cz/x", "title": "A"}, {"url": "https://a.cz/x"}]})
    assert w == [{"kind": "web", "ref": "https://a.cz/x", "label": "A", "tool": "web_search"}]
    m = intelplus.sources_from("mail_read", {}, {"from": "petr@x.cz", "from_name": "Petr", "subject": "Faktura"})
    assert m[0]["kind"] == "e-mail" and "Faktura" in m[0]["label"]
    assert intelplus.sources_from("remember", {"path": "a.txt"}, None) == []


async def test_agent_attaches_sources(rt, tmp_path):
    f = tmp_path / "pozn.txt"
    f.write_text("tajný recept: 3 vejce", encoding="utf-8")
    rt.router = FakeRouter([[("read_text_file", {"path": str(f)})], "Podle poznámky 3 vejce."])
    q = rt.bus.subscribe()
    await rt.agent.handle("kolik vajec?")
    done = None
    while not q.empty():
        m = q.get_nowait()
        if m["type"] == "assistant_done":
            done = m
    assert done and done["sources"][0]["kind"] == "soubor" and done["sources"][0]["label"] == "pozn.txt"
    last = rt.memory.history()[-1]
    assert last["meta"]["sources"][0]["ref"] == str(f)
