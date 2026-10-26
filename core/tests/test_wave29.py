"""OMEGA wave 29 (business, part 1): quote register with numbering and win rate, billing basis from a work log,
customer requests with priorities and due dates, supplier offers with weight sensitivity, campaign report, and income
/ expenses from a Fio-style bank statement (decimal commas, spaces in thousands)."""
import datetime as dt

import docx
import pytest

from jarvis import bizops, tools

tools.load_builtin()


def _dir(tmp_path, name):
    d = tmp_path / name                                                    # tmp_path itself also holds JARVIS_HOME
    d.mkdir()
    return d


async def test_quotes_register(rt):
    items = [{"name": "Web", "qty": 1, "price": 30000}]
    a = await tools.call("price_quote", {"client": "Pekárna Novák", "items": items, "valid_days": 30}, rt)
    b = await tools.call("price_quote", {"client": "Autoservis Dvořák", "items": [{"name": "E-shop", "qty": 1, "price": 80000}], "valid_days": 5}, rt)
    c = await tools.call("price_quote", {"client": "Kavárna U Mostu", "items": [{"name": "Logo", "qty": 1, "price": 9000}]}, rt)
    y = dt.date.today().year
    assert [a.data["number"], b.data["number"], c.data["number"]] == [f"N-{y}-0001", f"N-{y}-0002", f"N-{y}-0003"]
    assert f"N-{y}-0001" in "\n".join(p.text for p in docx.Document(a.data["path"]).paragraphs)
    assert (await tools.call("quote_status", {"which": "pekárna", "status": "přijatá"}, rt)).data["number"] == f"N-{y}-0001"
    assert (await tools.call("quote_status", {"which": f"N-{y}-0003", "status": "odmítnutá"}, rt)).ok
    assert not (await tools.call("quote_status", {"which": "nikdo", "status": "přijatá"}, rt)).ok
    r = (await tools.call("quotes_overview", {}, rt)).data
    assert r["open"] == 1 and r["open_value"] == 96800 and r["expiring"][0]["client"] == "Autoservis Dvořák" and r["win_rate"] == 50
    later = bizops.quotes_overview(rt, today=dt.date.today() + dt.timedelta(days=10))
    assert later["open"] == 0 and later["expired"] == 1


async def test_billing_basis(rt, tmp_path):
    d = _dir(tmp_path, "vykaz")
    (d / "vykaz.csv").write_text(
        "Datum;Klient;Projekt;Hodiny;Sazba;Popis;Náklady;Fakturováno\n"
        "2026-10-01;Pekárna Novák;Web;3,5;1200;úprava menu;;\n"
        "2026-10-02;Pekárna Novák;Web;2;1200;nové fotky;450;\n"
        "2026-10-03;Pekárna Novák;Údržba;1;900;aktualizace;;\n"
        "2026-09-28;Pekárna Novák;Web;5;1200;návrh;;ano\n"
        "2026-10-05;Autoservis Dvořák;E-shop;10;;napojení plateb;;\n", encoding="utf-8")
    r = await tools.call("billing_basis", {"file": str(d / "vykaz.csv"), "period": "2026-10"}, rt)
    nov = r.data["clients"]["Pekárna Novák"]
    assert nov["hours"] == 6.5 and nov["base"] == 5.5 * 1200 + 900 + 450 and nov["total"] == round(nov["base"] * 1.21, 2)
    assert r.data["clients"]["Autoservis Dvořák"]["missing_rate"] == 1 and "Bez sazby u Autoservis Dvořák" in r.summary
    assert len(r.attachments) == 2 and "Podklad k fakturaci – Pekárna Novák" in docx.Document(r.attachments[0]).paragraphs[0].text
    all_ = await tools.call("billing_basis", {"file": str(d / "vykaz.csv"), "client": "pekárna", "default_rate": 1000, "documents": False}, rt)
    assert all_.data["skipped_billed"] == 1 and list(all_.data["clients"]) == ["Pekárna Novák"]


async def test_customer_requests(rt):
    today = dt.date.today()
    a = await tools.call("customer_request", {"client": "Pekárna Novák", "text": "Na webu nefunguje objednávkový formulář!"}, rt)
    b = await tools.call("customer_request", {"client": "Kavárna", "text": "Až budete mít čas, šlo by přidat Instagram?"}, rt)
    await tools.call("customer_request", {"client": "Pekárna Novák", "text": "Prosím o změnu otevírací doby."}, rt)
    assert a.data["priority"] == "vysoká" and a.data["due"] == (today + dt.timedelta(days=1)).isoformat() and b.data["priority"] == "nízká"
    r = bizops.requests(rt, today=today + dt.timedelta(days=7))
    assert r["requests"][0]["id"] == a.data["id"] and r["requests"][0]["overdue_days"] == 6 and r["overdue"] == 2 and r["by_client"]["Pekárna Novák"] == 2
    await tools.call("customer_request", {"close": "formulář"}, rt)
    left = (await tools.call("customer_requests", {}, rt)).data
    assert left["open"] == 2 and all("formulář" not in x["text"] for x in left["requests"])
    assert not (await tools.call("customer_request", {"close": "neexistuje"}, rt)).ok


async def test_compare_offers(rt):
    offers = [{"supplier": "Alfa", "price": 100000, "delivery_days": 30, "warranty_months": 24, "rating": 4.5},
              {"supplier": "Beta", "price": 90000, "delivery_days": 60, "warranty_months": 12, "rating": 4.0},
              {"supplier": "Gama", "price": 120000, "delivery_days": 10, "warranty_months": 36, "rating": 4.8}]
    r = await tools.call("compare_offers", {"offers": offers, "weights": {"price": 0.6, "delivery_days": 0.2, "warranty_months": 0.2}}, rt)
    s = {x["supplier"]: x for x in r.data["ranking"]}
    assert s["Beta"]["parts"]["price"] == 1 and s["Gama"]["parts"]["price"] == 0 and s["Alfa"]["parts"]["price"] == pytest.approx(2 / 3, abs=1e-3)
    assert r.data["winner"] == "Alfa" and s["Alfa"]["score"] == pytest.approx(100 * (0.6 * 2 / 3 + 0.2 * 0.6 + 0.2 * 0.5), abs=0.1)
    assert not r.data["robust"] and r.data["flips"]
    tco = bizops.compare_offers([{"supplier": "Levná", "price": 10000, "yearly_cost": 5000}, {"supplier": "Úsporná", "price": 20000, "yearly_cost": 1000}],
                                {"price": 1}, years=5)
    assert tco["tco"] == {"Levná": 35000, "Úsporná": 25000} and tco["winner"] == "Úsporná"
    assert not (await tools.call("compare_offers", {"offers": offers[:1]}, rt)).ok


async def test_campaign_report(rt, tmp_path):
    d = _dir(tmp_path, "kampane")
    (d / "k.csv").write_text("Kampaň;Náklady;Zobrazení;Prokliky;Konverze;Tržby\n"
                             "Google brand;5 000;20 000;1 000;80;40 000\n"
                             "Facebook;8 000;200 000;2 000;20;6 000\n"
                             "Google brand;1 000;4 000;200;16;8 000\n", encoding="utf-8")
    r = await tools.call("campaign_report", {"file": str(d / "k.csv")}, rt)
    g = r.data["campaigns"][0]
    assert g["campaign"] == "Google brand" and g["spend"] == 6000 and g["roas"] == 8 and g["cpa"] == 62.5 and g["ctr_pct"] == 5 and g["cr_differs"]
    fb = r.data["campaigns"][1]
    assert fb["roas"] == 0.75 and fb["roi_pct"] == -25 and fb["cpa"] == 400
    assert r.data["move_budget"] == {"from": ["Facebook"], "to": "Google brand", "amount": 8000, "extra_revenue_if_roas_holds": 58000}
    assert r.data["total"]["roas"] == pytest.approx(54000 / 14000, abs=0.01)


async def test_cashflow_from_bank_statement(rt, tmp_path):
    d = _dir(tmp_path, "banka")
    rows = ["Datum;Objem;Měna;Protiúčet;Zpráva pro příjemce;Poznámka;Typ"]
    for m in (7, 8, 9, 10):
        rows += [f"15.{m:02d}.2026;45 000,00;CZK;123/0800;Mzda {m}/2026;MZDA FIRMA SRO;Příjem",
                 f"01.{m:02d}.2026;-15 500,00;CZK;555/0100;Nájem;Nájemné byt;Platba",
                 f"03.{m:02d}.2026;-299,00;CZK;;;Netflix.com;Platba kartou",
                 f"05.{m:02d}.2026;-1 840,50;CZK;;;Albert Praha;Platba kartou",
                 f"20.{m:02d}.2026;-2 100,00;CZK;;;Lidl dekuje za nakup;Platba kartou"]
    rows += ["21.10.2026;-12 990,00;CZK;;;ALZA.CZ notebook;Platba kartou", "22.10.2026;-640,00;CZK;;;Neznamy obchod XY;Platba kartou"]
    (d / "fio.csv").write_text("\n".join(rows) + "\n", encoding="cp1250")
    r = await tools.call("cashflow_report", {"file": str(d / "fio.csv")}, rt)
    oct_ = r.data["months"][-1]
    assert oct_["month"] == "2026-10" and oct_["income"] == 45000 and oct_["expenses"] == pytest.approx(15500 + 299 + 1840.5 + 2100 + 12990 + 640)
    cats = {c["category"]: c["amount"] for c in r.data["categories"]}
    assert cats["bydlení"] == 4 * 15500 and cats["jídlo"] == pytest.approx(4 * (1840.5 + 2100)) and cats["předplatné"] == 4 * 299 and cats["nákupy"] == 12990
    rec = {x["category"]: x["monthly"] for x in r.data["recurring"]}
    assert rec["bydlení"] == 15500 and rec["předplatné"] == 299
    assert r.data["spikes"] == [{"category": "nákupy", "last": 12990, "usual": 0}]
    assert r.data["largest"][0]["amount"] == -15500 and "Pravidelné platby" in r.summary
    own = bizops.cashflow(["Datum", "Částka", "Popis"], [["2026-10-01", "-500", "Fitko Pulse"]], {"sport": ["fitko"]})
    assert own["categories"][0]["category"] == "sport"
    with pytest.raises(ValueError):
        bizops.cashflow(["A", "B"], [["1", "2"]])


# ---------------------------------------------------------------------------------------------- part 2: web and customers
import httpx  # noqa: E402

from jarvis import webbiz  # noqa: E402

LONG = " ".join(["Pečeme kváskový chléb z české mouky každé ráno, rozvoz po Brně a okolí."] * 30)
SITE = {
    "/": f'<html lang="cs"><head><title>Pekárna Novák – kváskový chléb Brno</title><meta name="description" content="Kváskový chléb, '
         f'rohlíky a koláče pečené každé ráno. Rozvoz po Brně a objednávky online."><meta property="og:title" content="Pekárna"></head>'
         f'<body><h1>Pekárna Novák</h1><p>{LONG}</p><a href="/produkty">Produkty</a> <a href="/kontakt">Kontakt</a> <a href="/stara">Stará</a>'
         f'<a href="https://jinde.cz/">Jinde</a><img src="a.jpg" alt="chléb"></body></html>',
    "/produkty": '<html><head><title>Produkty</title></head><body><h1>Produkty</h1><h1>Další</h1><img src="b.jpg"><img src="c.jpg" alt="">'
                 '<p>Chléb, rohlík.</p><a href="/">Domů</a></body></html>',
    "/kontakt": '<html lang="cs"><head><title>Produkty</title><meta name="robots" content="noindex"></head><body><h1>Kontakt</h1><p>Tel.</p></body></html>',
}


def site(req: httpx.Request) -> httpx.Response:
    path = req.url.path
    if req.url.host == "pekarna.cz":
        if path == "/robots.txt":
            return httpx.Response(200, text="User-agent: *\nDisallow: /admin")
        if path == "/sitemap.xml":
            return httpx.Response(404)
        if path in SITE:
            return httpx.Response(200, text=SITE[path], headers={"content-type": "text/html; charset=utf-8"})
        return httpx.Response(404, text="nenalezeno")
    if req.url.host == "konkurence.cz":
        return httpx.Response(200, headers={"content-type": "text/html"}, text=COMP[0])
    return httpx.Response(404)


COMP = ['<html><head><title>Pekárna Konkurent</title></head><body><h1>Bezlepkový chléb a kváskový chléb</h1>'
        '<h2>Bezlepkové pečivo</h2><p>Bezlepkový chléb pro celiaky. Bezlepkový chléb každý den, rozvoz zdarma, rozvoz zdarma nad 500 Kč.</p>'
        '<p>Kváskový chléb</p><p>89 Kč</p><p>Bezlepková bageta</p><p>65 Kč</p></body></html>']


async def test_seo_audit(rt):
    rt.transport = httpx.MockTransport(site)
    r = await tools.call("seo_audit", {"url": "pekarna.cz"}, rt)
    pages = {p["url"].split("pekarna.cz")[1]: p for p in r.data["pages"]}
    assert set(pages) == {"/", "/produkty", "/kontakt"} and pages["/"]["problems"] == []
    pr = pages["/produkty"]["problems"]
    assert "chybí meta description" in pr and "2× H1 (má být právě jeden)" in pr and "2 obrázků bez alt textu" in pr and any("málo textu" in x for x in pr)
    assert "stránka je noindex – Google ji nezobrazí" in pages["/kontakt"]["problems"]
    assert r.data["broken"] == [{"url": "https://pekarna.cz/stara", "status": 404}] and not r.data["sitemap"] and r.data["robots"]
    assert any("stejný titulek" in i for i in r.data["site_issues"]) and "chybí sitemap.xml" in r.summary and 0 < r.data["score"] < 100


async def test_keyword_gap_and_competitor_watch(rt):
    rt.transport = httpx.MockTransport(site)
    g = await tools.call("keyword_gap", {"my_url": "https://pekarna.cz/", "competitor_url": "https://konkurence.cz/"}, rt)
    terms = [x["term"] for x in g.data["gaps"]]
    assert "bezlepkovy" in terms and "bezlepkovy chleb" in terms and g.data["gaps"][0]["in_headings"] and "chleb" not in terms
    first = await tools.call("competitor_watch", {"url": "konkurence.cz"}, rt)
    assert first.data["first"] and first.data["now"]["prices"] == {"Kváskový chléb": 89, "Bezlepková bageta": 65}
    COMP[0] = COMP[0].replace("89 Kč", "99 Kč").replace("<p>Bezlepková bageta</p><p>65 Kč</p>", "<p>Špaldový rohlík</p><p>1 290,50 Kč</p>")
    second = await tools.call("competitor_watch", {"url": "konkurence.cz"}, rt)
    d = second.data
    assert d["price_changes"] == [{"item": "Kváskový chléb", "before": 89, "after": 99, "pct": 11.2}]
    assert d["new_items"] == ["Špaldový rohlík"] and d["removed_items"] == ["Bezlepková bageta"] and d["now"]["prices"]["Špaldový rohlík"] == 1290.5
    m = await tools.call("feature_matrix", {"urls": ["https://pekarna.cz/", "https://konkurence.cz/"],
                                            "features": ["rozvoz", "bezlepkové / bezlepkový", "koláče", "e-shop"]}, rt)
    mx = m.data["matrix"]
    assert mx["rozvoz"] == {"https://pekarna.cz/": True, "https://konkurence.cz/": True}
    assert mx["bezlepkové / bezlepkový"]["https://pekarna.cz/"] is False and mx["e-shop"]["https://konkurence.cz/"] is False
    assert m.data["unique"]["https://konkurence.cz/"] == ["bezlepkové / bezlepkový"]


async def test_bought_together(rt, tmp_path):
    d = _dir(tmp_path, "obj")
    rows = ["objednávka;zákazník;produkt"]
    for i in range(10):
        rows += [f"{i};z{i}@x.cz;Kávovar", f"{i};z{i}@x.cz;Odvápňovač"] if i < 6 else [f"{i};z{i}@x.cz;Kávovar"]
        if i % 2:
            rows.append(f"{i};z{i}@x.cz;Hrnek")
    rows += ["20;novak@x.cz;Odvápňovač", "21;novak@x.cz;Káva", "22;dvorak@x.cz;Káva", "23;dvorak@x.cz;Hrnek"]
    (d / "o.csv").write_text("\n".join(rows) + "\n", encoding="utf-8")
    r = await tools.call("bought_together", {"file": str(d / "o.csv")}, rt)
    top = r.data["rules"][0]
    assert {top["if"], top["then"]} == {"Kávovar", "Odvápňovač"} and top["together"] == 6
    k2o = next(x for x in r.data["rules"] if x["if"] == "Kávovar" and x["then"] == "Odvápňovač")
    assert k2o["confidence_pct"] == 60 and k2o["lift"] == pytest.approx(6 * 14 / (10 * 7), abs=0.01)
    cs = next(t for t in r.data["cross_sell"] if t["offer"] == "Odvápňovač")
    assert cs["customers"] == ["z6@x.cz", "z7@x.cz", "z8@x.cz", "z9@x.cz"]


async def test_follow_ups_and_dashboard(rt, tmp_path):
    today = dt.date.today()
    bizops.register_quote(rt, "N-2026-0101", "Stará firma", 1000, 1210, 30, "", today=today - dt.timedelta(days=10))
    bizops.register_quote(rt, "N-2026-0102", "Spěchá s.r.o.", 5000, 6050, 2, "", today=today)
    bizops.register_quote(rt, "N-2026-0103", "Věrný klient", 9000, 10890, 30, "", today=today - dt.timedelta(days=200))
    rt.db.exec("UPDATE quotes SET status='přijatá', decided=? WHERE number='N-2026-0103'", (__import__("time").time() - 120 * 86400,))
    bizops.add_request(rt, "Pekárna", "Nefunguje e-shop", today=today - dt.timedelta(days=5))
    r = await tools.call("follow_ups", {}, rt)
    whys = [(x["who"], x["why"]) for x in r.data]
    assert whys[0][0] == "Pekárna" and "po termínu" in whys[0][1]
    assert ("Spěchá s.r.o.", "nabídka N-2026-0102 propadne za 2 d") in whys and any(w[0] == "Stará firma" and "bez odpovědi 10 dní" in w[1] for w in whys)
    assert any(w[0] == "Věrný klient" and "bez nové zakázky" in w[1] for w in whys) and all(x["draft"] for x in r.data)
    d = _dir(tmp_path, "dash")
    (d / "b.csv").write_text("Datum;Částka;Popis\n2026-08-05;30 000;faktura\n2026-08-10;-10 000;nájem\n2026-09-05;20 000;faktura\n"
                             "2026-09-10;-15 000;nájem\n2026-10-05;10 000;faktura\n2026-10-10;-15 000;nájem\n", encoding="utf-8")
    dash = await tools.call("business_dashboard", {"bank_file": str(d / "b.csv")}, rt)
    k = dash.data["kpi"]
    assert k["otevřené nabídky"] == 2 and k["požadavky po termínu"] == 1 and k["čistý výsledek"] == -5000 and k["lidé k oslovení"] == len(r.data)
    assert any("ve ztrátě" in a for a in dash.data["alerts"]) and any("klesá" in a for a in dash.data["alerts"])


def test_snapshot_prices():
    s = webbiz.snapshot("<p>Dort</p><p>1.290 Kč</p><p>Káva</p><p>49,- </p><p>Bez ceny</p>")
    assert s["prices"] == {"Dort": 1290, "Káva": 49}
